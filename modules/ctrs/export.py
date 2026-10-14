"""A CTR as an Excel workbook laid out like the company's CTR sheets (header block, scope of work, resources with their
hours and cost, additions, payment terms, basis of proposal), so it can be edited by hand afterwards.

The printed part is columns A to J. Beside the tables, outside the print area, sit the working figures the original sheets
keep there: the factor of each scope line, the two norms (columns K, L, M), and for each resource the percentage split and
where its hours come from. Line hours, resource costs and every total are formulas. A resource's hours are the figure
this tool shared out in whole hours (calc.share_out), written as a value, so change a split here and in the tool, not by
editing the sheet.

ONE sheet, A4 portrait, full width, with a page break after the scope table, as the original sheets have.
"""
from __future__ import annotations

import math
import re
import zipfile
from datetime import date
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import openpyxl
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.pagebreak import Break
from openpyxl.worksheet.properties import PageSetupProperties

from . import calc

LOGO = Path(__file__).parent / "assets" / "nmc-logo.png"
WIDTHS = {"A": 22.29, "B": 5.57, "C": 9.29, "D": 10.0, "E": 8.43, "F": 11.57, "G": 16.29, "H": 8.43, "I": 14.29, "J": 14.29,
          "K": 9.0, "L": 12.0, "M": 12.0}
LAST = 10                                       # column J: the last printed column
FONT, SIZE = "Aptos", 10
MIN_RESOURCE_ROWS, MIN_ADDITION_ROWS = 4, 2     # the sheets always show at least this many
FULL_CHARS, SCOPE_CHARS = 135, 60               # about how much fits on one line of A:J and of A:E
LINE_HEIGHT = 13.2
_MED, _THIN = Side(style="medium"), Side(style="thin")
_BLACK = PatternFill("solid", fgColor="000000")
_GREEN = PatternFill("solid", fgColor="92D050")
_GREY = PatternFill("solid", fgColor="BFBFBF")
_GBP = "[$£-809]#,##0.00"
_HOURS = "0"


def text(value) -> str:
    """Text fit for a cell: nothing for None, and no control characters (which Excel files cannot hold)."""
    return "" if value is None else ILLEGAL_CHARACTERS_RE.sub("", str(value))


def lines_needed(value: str, chars: int) -> int:
    return sum(max(1, math.ceil(len(part) / chars)) for part in text(value).split("\n"))


def put(ws, row: int, first: int, last: int, value=None, *, black=False, fill=None, bold=False, align="left",
        valign="center", fmt: Optional[str] = None, size=SIZE, italic=False, thin=False, formula=False) -> None:
    """Write into columns first..last of a row (merged when more than one column), boxed all round: a medium line for the
    header blocks, a thin one inside the scope table. `black` is the white-on-black label style. A value starting with =
    is typed text and stays text, unless `formula` is set."""
    if last > first:
        ws.merge_cells(start_row=row, start_column=first, end_row=row, end_column=last)
    side = _THIN if thin else _MED
    for col in range(first, last + 1):
        cell = ws.cell(row, col)
        cell.border = Border(left=side, right=side, top=side, bottom=side)
        cell.font = Font(name=FONT, size=size, bold=bold or black, italic=italic, color="FFFFFF" if black else None)
        cell.alignment = Alignment(horizontal=align, vertical=valign, wrap_text=True)
        if black:
            cell.fill = _BLACK
        elif fill:
            cell.fill = fill
    cell = ws.cell(row, first)
    if value is not None and value != "":
        cell.value = value
        if isinstance(value, str) and value.startswith("=") and not formula:
            cell.data_type = "s"                       # typed text that starts with = stays text, never a formula
    if fmt:
        cell.number_format = fmt


def side_note(ws, row: int, col: int, value, *, fmt: Optional[str] = None, bold=False) -> None:
    """A working figure outside the print area (columns K to M): plain, no box."""
    cell = ws.cell(row, col)
    if value is not None and value != "":
        cell.value = value
    cell.font = Font(name=FONT, size=SIZE, bold=bold, color="595959")
    cell.alignment = Alignment(horizontal="left", vertical="center")
    if fmt:
        cell.number_format = fmt


def _date(value: str):
    try:
        return date.fromisoformat(text(value).strip()[:10])
    except ValueError:
        return None


def _runs(rows: List[int]) -> List[Tuple[int, int]]:
    """Consecutive row numbers as (first, last) ranges."""
    out: List[Tuple[int, int]] = []
    for r in rows:
        if out and out[-1][1] == r - 1:
            out[-1] = (out[-1][0], r)
        else:
            out.append((r, r))
    return out


def _block(ws, r: int, title: str, body: str, minimum: float) -> int:
    """A black title band and, under it, a block of text across the page. Returns the next free row."""
    put(ws, r, 1, LAST, title, black=True, align="center")
    put(ws, r + 1, 1, LAST, text(body), valign="top")
    ws.row_dimensions[r + 1].height = max(minimum, LINE_HEIGHT * lines_needed(body, FULL_CHARS) + 6)
    return r + 2


def _header(ws, ctr: Dict) -> int:
    for r in range(1, 5):
        ws.row_dimensions[r].height = 15
    ws.merge_cells("A1:D4")                      # the logo sits in A1:D4
    ws.merge_cells("E1:J4")
    for r in range(1, 5):
        for c in range(1, LAST + 1):
            ws.cell(r, c).border = Border(left=_MED if c == 1 else None, right=_MED if c == LAST else None,
                                          top=_MED if r == 1 else None, bottom=_MED if r == 4 else None)
    for r, (label, key) in enumerate([("Project Name", "project_name"), ("Project Title", "project_title"),
                                      ("Client", "client"), ("Date", "ctr_date")], 5):
        put(ws, r, 1, 1, label, black=True)
        when = _date(ctr.get(key)) if key == "ctr_date" else None
        put(ws, r, 2, LAST, when or text(ctr.get(key)), bold=key == "ctr_date", fmt="dd/mm/yyyy" if when else None)
    put(ws, 9, 1, 2, "Framework Agreement Ref", black=True); put(ws, 9, 3, LAST, text(ctr.get("framework_ref")))
    put(ws, 10, 1, LAST); ws.row_dimensions[10].height = 6
    put(ws, 11, 1, 1, "CTR Title", black=True); put(ws, 11, 2, 5, text(ctr.get("title")))
    put(ws, 11, 6, 6); put(ws, 11, 7, 7)
    put(ws, 11, 8, 8, "CTR No", black=True); put(ws, 11, 9, LAST, text(ctr.get("ctr_no")), align="center")
    put(ws, 12, 1, LAST); ws.row_dimensions[12].height = 6
    put(ws, 13, 1, 2, "Prepared by", black=True, align="center"); put(ws, 13, 3, 5, "Checked by", black=True, align="center")
    put(ws, 13, 6, 7, black=True); put(ws, 13, 8, LAST, "Approved by", black=True, align="center")
    put(ws, 14, 1, 2, text(ctr.get("prepared_by")), align="center"); put(ws, 14, 3, 5, text(ctr.get("checked_by")), align="center")
    put(ws, 14, 6, 7); put(ws, 14, 8, LAST, text(ctr.get("approved_by")), align="center")
    r = _block(ws, 15, "Objectives", ctr.get("objectives"), 45)
    r = _block(ws, r, "Input Requirements", ctr.get("inputs"), 45)
    return r


def _scope(ws, ctr: Dict, calced: Dict, top: int, results: Dict[str, float]) -> int:
    """The Scope of Work table. Returns the last row used."""
    p1, p2 = calced["labels"]
    put(ws, top, 1, LAST, "Scope of Work", black=True, align="center")
    head = top + 1
    put(ws, head, 1, 5, text(ctr.get("scope_intro")), bold=True, thin=True)
    for col, label, al in ((6, "Activity/ Deliverable", "left"), (7, "WBS", "left"), (8, "Qty", "right"),
                           (9, p1, "right"), (10, p2, "right")):
        put(ws, head, col, col, text(label), bold=True, thin=True, align=al)
    ws.row_dimensions[head].height = 26.2
    side_note(ws, head, 11, "Factor", bold=True); side_note(ws, head, 12, f"{p1} NORM", bold=True)
    side_note(ws, head, 13, f"{p2} NORM", bold=True)
    lines = ctr.get("lines", [])
    first = head + 1
    plain = [first + i for i, ln in enumerate(lines) if ln.get("kind") != "allowance"]
    spans = ",".join(f"I{a}:J{b}" for a, b in _runs(plain))
    for i, ln in enumerate(lines):
        r, res = first + i, calced["lines"][i]
        allowance = ln.get("kind") == "allowance"
        put(ws, r, 1, 5, text(ln.get("description")), thin=True)
        put(ws, r, 6, 6, text(ln.get("flag")), thin=True)
        put(ws, r, 7, 7, text(ln.get("wbs")), thin=True)
        if allowance:
            pct = ln.get("qty")
            put(ws, r, 8, 8, None if pct is None else pct / 100.0, thin=True, align="right",
                fmt="0%" if pct is None or float(pct).is_integer() else "0.0#%")
        else:
            put(ws, r, 8, 8, ln.get("qty"), thin=True, align="right")
        for col, key, norm in ((9, "eng_override", "L"), (10, "des_override", "M")):
            value = res["eng" if col == 9 else "des"]
            if ln.get(key) is not None:                  # hours typed on the line: a plain number
                put(ws, r, col, col, value, thin=True, align="right", fmt=_HOURS)
            elif allowance and col == 10:                # an allowance only adds to the first column
                put(ws, r, col, col, 0, thin=True, align="right", fmt=_HOURS)
            else:
                if allowance:                            # a percentage of every other line (not of other allowances)
                    fmla = f"=ROUNDUP(ROUND(SUM({spans})*H{r},6),0)" if spans else "=0"
                else:
                    fmla = f'=IF(OR(H{r}="",{norm}{r}=""),0,ROUNDUP(ROUND(H{r}*{norm}{r}*IF(K{r}="",1,K{r}),6),0))'
                put(ws, r, col, col, fmla, thin=True, align="right", fmt=_HOURS, formula=True)
                results[f"{get_column_letter(col)}{r}"] = value
        if not allowance:
            side_note(ws, r, 11, ln.get("factor")); side_note(ws, r, 12, ln.get("eng_norm")); side_note(ws, r, 13, ln.get("des_norm"))
        ws.row_dimensions[r].height = max(15.0, LINE_HEIGHT * lines_needed(ln.get("description"), SCOPE_CHARS) + 2)
    last = first + max(len(lines), 1) - 1
    if not lines:                                        # an empty table still shows one row to fill in
        put(ws, first, 1, 5, thin=True)
        for col in range(6, LAST + 1):
            put(ws, first, col, col, thin=True)
    if ctr.get("source_text"):
        last += 1
        put(ws, last, 1, LAST, f"Hours from the jobcards: {text(ctr['source_text'])}", italic=True, thin=True)
    return last


def _after_scope(ws, ctr: Dict, calced: Dict, top: int, results: Dict[str, float]) -> int:
    """Notes, dates, resources, additions, payment terms and basis. Returns the last row used."""
    r = top
    for key in ("scope_notes", "assumptions", "exclusions"):
        body = text(ctr.get(key))
        if body.strip():
            put(ws, r, 1, LAST, body, valign="top")
            ws.row_dimensions[r].height = LINE_HEIGHT * lines_needed(body, FULL_CHARS) + 6
            r += 1
    ws.row_dimensions[r].height = 12
    r += 1
    put(ws, r, 1, 1, "Commencement Date", black=True); put(ws, r, 2, 5, text(ctr.get("commencement")), align="center")
    put(ws, r, 6, 6); put(ws, r, 7, 7)
    put(ws, r, 8, 8, "Duration", black=True); put(ws, r, 9, LAST, text(ctr.get("duration")), align="center")
    put(ws, r + 1, 1, 1, "Completion Date", black=True); put(ws, r + 1, 2, 5, text(ctr.get("completion")), align="center")
    put(ws, r + 1, 6, 6); put(ws, r + 1, 7, 7)
    put(ws, r + 1, 8, 8, "Expenses", black=True); put(ws, r + 1, 9, LAST, text(ctr.get("expenses")), align="center")
    ws.row_dimensions[r].height = ws.row_dimensions[r + 1].height = 15
    r += 2
    put(ws, r, 1, LAST); ws.row_dimensions[r].height = 6
    r += 1

    put(ws, r, 1, 2, "Resource", black=True, align="center"); put(ws, r, 3, 4, "Cost £ / Hour", black=True, align="center")
    put(ws, r, 5, 8, "Hours", black=True, align="center"); put(ws, r, 9, LAST, "Total", black=True, align="center")
    side_note(ws, r, 11, "Split", bold=True); side_note(ws, r, 12, "Takes hours from", bold=True)
    resources = calced["resources"]
    first = r + 1
    for i in range(max(MIN_RESOURCE_ROWS, len(resources))):
        row = first + i
        ws.row_dimensions[row].height = 15
        res = resources[i] if i < len(resources) else None
        put(ws, row, 1, 2, text(res["name"]) if res else None)
        put(ws, row, 3, 4, res["rate"] if res else None, align="center", fmt=_GBP)
        put(ws, row, 5, 8, res["hours"] if res else None, align="center", fmt=_HOURS)
        if res:
            put(ws, row, 9, LAST, f"=E{row}*C{row}", align="center", fmt=_GBP, formula=True)
            results[f"I{row}"] = res["cost"]
            if res["source"] != calc.FIXED:
                side_note(ws, row, 11, (res["split_pct"] if res["split_pct"] is not None else 100) / 100.0, fmt="0.##%")
                side_note(ws, row, 12, calc.source_label(res["source"], calced["labels"]))
            else:
                side_note(ws, row, 12, "typed hours")
        else:
            put(ws, row, 9, LAST, None, align="center", fmt=_GBP)
    last = first + max(MIN_RESOURCE_ROWS, len(resources)) - 1
    total_row = last + 1
    put(ws, total_row, 1, 4, black=True)
    put(ws, total_row, 5, 8, f"=SUM(E{first}:E{last})", fill=_GREEN, bold=True, align="center", fmt=_HOURS, formula=True)
    put(ws, total_row, 9, LAST, f"=SUM(I{first}:I{last})", bold=True, align="center", fmt=_GBP, formula=True)
    results[f"E{total_row}"], results[f"I{total_row}"] = calced["totals"]["hours"], calced["totals"]["resources_cost"]

    r = total_row + 1
    put(ws, r, 1, 2, "Addition", black=True, align="center"); put(ws, r, 3, 4, "Cost £ / Unit", black=True, align="center")
    put(ws, r, 5, 8, "Unit", black=True, align="center"); put(ws, r, 9, LAST, "Total", black=True, align="center")
    additions = calced["additions"]
    first_add = r + 1
    for i in range(max(MIN_ADDITION_ROWS, len(additions))):
        row = first_add + i
        ws.row_dimensions[row].height = 15
        add = additions[i] if i < len(additions) else None
        put(ws, row, 1, 2, text(add["name"]) if add else None)
        put(ws, row, 3, 4, add["unit_cost"] if add else None, align="center", fmt=_GBP)
        put(ws, row, 5, 8, add["units"] if add else None, align="center", fmt="General")
        put(ws, row, 9, LAST, f"=C{row}*E{row}" if add else None, fill=_GREY, align="center", fmt=_GBP, formula=True)
        if add:
            results[f"I{row}"] = add["cost"]
    last_add = first_add + max(MIN_ADDITION_ROWS, len(additions)) - 1
    grand = last_add + 1
    put(ws, grand, 1, 8, black=True)
    put(ws, grand, 9, LAST, f"=SUM(I{first_add}:I{last_add})+I{total_row}", fill=_GREEN, bold=True, align="center",
        fmt=_GBP, formula=True)
    results[f"I{grand}"] = calced["totals"]["cost"]

    r = grand + 1
    ws.row_dimensions[r].height = 15
    terms = ctr.get("payment_terms")
    put(ws, r, 1, 1, "Payment Terms", black=True)
    put(ws, r, 2, 3, "Reimbursable", black=True, align="center"); put(ws, r, 4, 4, "X" if terms == "Reimbursable" else None, align="center")
    put(ws, r, 5, 7, "Budget", black=True); put(ws, r, 8, 8, "X" if terms == "Budget" else None, align="center")
    put(ws, r, 9, 9, "Lump Sum", black=True); put(ws, r, 10, 10, "X" if terms == "Lump Sum" else None, align="center")
    r += 1
    put(ws, r, 1, LAST); ws.row_dimensions[r].height = 6
    return _block(ws, r + 1, "Basis of Proposal", ctr.get("basis"), 40) - 1


def _setup(ws, last_row: int, break_after: int) -> None:
    for letter, width in WIDTHS.items():
        ws.column_dimensions[letter].width = width
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_margins.left = ws.page_margins.right = 0.7
    ws.page_margins.top = ws.page_margins.bottom = 0.75
    ws.print_area = f"A1:J{last_row}"
    ws.row_breaks.append(Break(id=break_after))


def _box(ws, last_row: int) -> None:
    """The heavy line round the whole printed area."""
    for r in range(1, last_row + 1):
        for c, side in ((1, "left"), (LAST, "right")):
            cell = ws.cell(r, c)
            b = cell.border
            cell.border = Border(left=_MED if side == "left" else b.left, right=_MED if side == "right" else b.right,
                                 top=b.top, bottom=b.bottom)


# ---- the finished file: cached results and the logo, added to the zip -----------------------------------
_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_DRAWING = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
    '<xdr:oneCellAnchor><xdr:from><xdr:col>0</xdr:col><xdr:colOff>272415</xdr:colOff><xdr:row>0</xdr:row>'
    '<xdr:rowOff>100965</xdr:rowOff></xdr:from><xdr:ext cx="1082675" cy="520065"/>'
    '<xdr:pic><xdr:nvPicPr><xdr:cNvPr id="2" name="Logo" descr="NMC Energy"/><xdr:cNvPicPr>'
    '<a:picLocks noChangeAspect="1"/></xdr:cNvPicPr></xdr:nvPicPr>'
    f'<xdr:blipFill><a:blip xmlns:r="{_REL}" r:embed="rId1"/><a:stretch><a:fillRect/></a:stretch></xdr:blipFill>'
    '<xdr:spPr><a:xfrm><a:off x="272415" y="100965"/><a:ext cx="1082675" cy="520065"/></a:xfrm>'
    '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></xdr:spPr></xdr:pic><xdr:clientData/></xdr:oneCellAnchor></xdr:wsDr>')
_DRAWING_RELS = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                 f'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                 f'<Relationship Id="rId1" Type="{_REL}/image" Target="../media/image1.png"/></Relationships>')
_SHEET_RELS = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
               '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
               f'<Relationship Id="rId1" Type="{_REL}/drawing" Target="../drawings/drawing1.xml"/></Relationships>')


def _finish(path: Path, results: Dict[str, float], logo: Optional[bytes]) -> None:
    """Put each formula's result in the file beside the formula (openpyxl writes only the formula, and a viewer that does
    not calculate, such as Excel's Protected View, would show the totals empty), and add the logo picture."""
    sheet, temp = "xl/worksheets/sheet1.xml", path.with_name(path.name + ".tmp")
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(temp, "w", zipfile.ZIP_DEFLATED) as target:
        if "xl/worksheets/_rels/sheet1.xml.rels" in source.namelist():
            logo = None                                  # the sheet already has parts of its own; skip the picture, not the file
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == sheet:
                xml = data.decode("utf-8")
                for ref, value in results.items():
                    xml = re.sub(rf'(<c r="{ref}"[^>]*><f>[^<]*</f>)(?:<v></v>|<v\s*/>)',
                                 lambda m: f"{m.group(1)}<v>{value:.10g}</v>", xml, count=1)
                if logo:
                    xml = xml.replace("</worksheet>", f'<drawing xmlns:r="{_REL}" r:id="rId1"/></worksheet>')
                data = xml.encode("utf-8")
            elif item.filename == "[Content_Types].xml" and logo:
                xml = data.decode("utf-8")
                if 'Extension="png"' not in xml:
                    xml = xml.replace("<Override", '<Default Extension="png" ContentType="image/png"/><Override', 1)
                xml = xml.replace("</Types>", '<Override PartName="/xl/drawings/drawing1.xml" ContentType='
                                  '"application/vnd.openxmlformats-officedocument.drawing+xml"/></Types>')
                data = xml.encode("utf-8")
            target.writestr(item, data)
        if logo:
            target.writestr("xl/media/image1.png", logo)
            target.writestr("xl/drawings/drawing1.xml", _DRAWING)
            target.writestr("xl/drawings/_rels/drawing1.xml.rels", _DRAWING_RELS)
            target.writestr("xl/worksheets/_rels/sheet1.xml.rels", _SHEET_RELS)
    temp.replace(path)


def sheet_name(title: str) -> str:
    name = re.sub(r"[\[\]:*?/\\]", " ", text(title)).strip()[:31].strip()
    return name or "CTR"


def write_ctr(path: Path, ctr: Dict, calced: Optional[Dict] = None) -> None:
    """Write one CTR (as Repo.get_ctr gives it) to an .xlsx file. `calced` is calc.compute(ctr), worked out here if not given."""
    calced = calced or calc.compute(ctr)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name(ctr.get("title"))
    results: Dict[str, float] = {}
    top = _header(ws, ctr)
    scope_end = _scope(ws, ctr, calced, top, results)
    last = _after_scope(ws, ctr, calced, scope_end + 1, results)
    _box(ws, last)
    _setup(ws, last, scope_end)
    wb.properties.title = text(ctr.get("title"))
    wb.properties.creator = "NMC Tools"
    wb.save(path)
    _finish(path, results, LOGO.read_bytes() if LOGO.exists() else None)
