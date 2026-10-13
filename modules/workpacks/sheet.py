"""Turns a saved workpack into the input workbook the Workpack Tools engine reads.

The engine is kept exactly as released, so instead of changing it we give it what it already understands: its own
blank input sheet (made by the engine itself), filled in from the database. Cells are found by their header text.
"""
from __future__ import annotations

from pathlib import Path
from typing import Dict, List

import openpyxl

from .engine import fill_workpack as F
from .fields import CERTIFICATE_FIELDS, ROW_KINDS, WORKPACK_FIELDS

TABLE_SHEETS = {"line": "Lines", "joint": "Joints", "tag": "Tags"}


def _value(text):
    """Blank becomes an empty cell. The engine treats an empty cell as 'leave blank on the form'."""
    return text if text not in (None, "") else None


def _put(cell, text) -> None:
    cell.value = _value(text)
    if isinstance(cell.value, str) and cell.value.startswith("="):
        cell.data_type = "s"            # text that starts with = must stay text, not become a formula


def _write_table(ws, records: List[Dict]) -> None:
    headers = {str(ws.cell(1, c).value).strip(): c for c in range(1, ws.max_column + 1) if ws.cell(1, c).value}
    for row, record in enumerate(records, 3):        # row 1 = headers, row 2 = explanations
        for label, text in record.items():
            _put(ws.cell(row, headers[label]), text)  # an unknown header is a bug, so let it fail loudly


def write_input_workbook(path: Path, workpack: Dict, certificates: List[Dict], rows: Dict[int, Dict[str, List[Dict]]]) -> None:
    """certificates are certificate dicts; rows is {certificate id: {row kind: [row dicts]}}."""
    F.make_sheet(str(path))
    wb = openpyxl.load_workbook(path)

    project = wb["Project"]
    where = {project.cell(r, 1).value: r for r in range(2, project.max_row + 1)}
    for column, label in WORKPACK_FIELDS:
        _put(project.cell(where[label], 2), workpack.get(column))

    _write_table(wb["Certificates"], [
        {"ID": c["cert_id"], "Form": c["form"], **{label: c[column] for column, label in CERTIFICATE_FIELDS}}
        for c in certificates])

    for kind, sheet in TABLE_SHEETS.items():
        columns = ROW_KINDS[kind][1]
        _write_table(wb[sheet], [
            {"Certificate ID": c["cert_id"], **{label: row[column] for column, label in columns}}
            for c in certificates for row in rows.get(c["id"], {}).get(kind, [])])
    wb.save(path)