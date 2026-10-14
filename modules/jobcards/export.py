"""A jobcard as an Excel workbook, laid out like the Construction Jobcard sheet (the print view), for the fill pack.

The workbook is meant to be edited by hand afterwards, so hours are formulas (quantity x norm x factor, and a total) and
every value sits in an ordinary cell. It is ONE sheet, like the original Excel jobcard: page 1 (details, flags, drawings,
completion), a page break, then page 2 (the header again and the tasks). One sheet matters because Excel's Print and
Save as PDF only do the active sheet, so a second sheet would never be printed. It is set up for A4 portrait, full width.
Task notes keep their bold, italic, underline and bullets as Excel rich text.
"""
from __future__ import annotations

import math
import re
import zipfile
from html.parser import HTMLParser
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import openpyxl
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.pagebreak import Break
from openpyxl.worksheet.properties import PageSetupProperties

from . import calc
from .constants import FLAG_GROUPS

MIN_TASK_ROWS = 10                                    # the sheet always has at least this many item rows
# Columns A to K, in the proportions of the print view. Together they are about as wide as an A4 portrait page with
# 0.4 inch margins (717 pixels), so Excel prints at full size instead of shrinking the sheet to fit.
WIDTHS = [8.25, 8.25, 7.5, 8.25, 13, 10.5, 6.75, 8.25, 7.5, 8, 7.25]
A4_PIXELS = 717
DESCRIPTION_CHARS = 58                                # about how much of a description fits on one line of B:G
INFO_CHARS = 110                                      # and of the additional information across A:K
FONT, SIZE = "Calibri", 10
_THIN = Side(style="thin")
_BOX = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_BLUE = PatternFill("solid", fgColor="DCE6F1")
_HIDE_ZERO = "General;-General;;@"


def text(value) -> str:
    """Text fit for a cell: nothing for None, and no control characters (which Excel files cannot hold)."""
    return "" if value is None else ILLEGAL_CHARACTERS_RE.sub("", str(value))


def put(ws, row: int, first: int, last: int, value=None, *, bold=False, blue=False, align="left", valign="center",
        size=SIZE, italic=False, fmt: Optional[str] = None, formula: bool = False) -> None:
    """Write a value into columns first..last of a row, merged when that is more than one column, boxed all round.
    A value starting with = is typed text and stays text, unless formula is set."""
    if last > first:
        ws.merge_cells(start_row=row, start_column=first, end_row=row, end_column=last)
    for col in range(first, last + 1):
        cell = ws.cell(row, col)
        cell.border = _BOX
        cell.font = Font(name=FONT, size=size, bold=bold, italic=italic)
        cell.alignment = Alignment(horizontal=align, vertical=valign, wrap_text=True)
        if blue:
            cell.fill = _BLUE
    cell = ws.cell(row, first)
    if value is not None:
        cell.value = value
        if isinstance(value, str) and value.startswith("=") and not formula:
            cell.data_type = "s"                       # typed text that starts with = stays text, never a formula
    if fmt:
        cell.number_format = fmt


# ---- notes as rich text ----------------------------------------------------------------------------
class _Runs(HTMLParser):
    """A cleaned note (see richtext.py) as runs of (text, bold, italic, underline), with bullets and line breaks."""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.runs: List[Tuple[str, bool, bool, bool]] = []
        self.on = {"b": 0, "i": 0, "u": 0}
        self.lists: List[List] = []

    def _break(self):
        if self.runs and not self.runs[-1][0].endswith("\n"):
            self.runs.append(("\n", False, False, False))

    def _add(self, value):
        self.runs.append((value, self.on["b"] > 0, self.on["i"] > 0, self.on["u"] > 0))

    def handle_starttag(self, tag, attrs):
        key = {"strong": "b", "em": "i"}.get(tag, tag)
        if key in self.on:
            self.on[key] += 1
        elif tag in ("ul", "ol"):
            self._break()
            self.lists.append([tag, 0])
        elif tag == "li":
            self._break()
            if self.lists and self.lists[-1][0] == "ol":
                self.lists[-1][1] += 1
                self._add(f"{self.lists[-1][1]}. ")
            else:
                self._add("• ")
        elif tag in ("div", "p"):
            self._break()
        elif tag == "br":
            self.runs.append(("\n", False, False, False))

    def handle_endtag(self, tag):
        key = {"strong": "b", "em": "i"}.get(tag, tag)
        if key in self.on:
            self.on[key] = max(0, self.on[key] - 1)
        elif tag in ("ul", "ol"):
            if self.lists:
                self.lists.pop()
            self._break()
        elif tag in ("li", "div", "p"):
            self._break()

    def handle_data(self, data):
        data = re.sub(r"\s+", " ", data)
        if data.strip() or (data and self.runs and not self.runs[-1][0].endswith("\n")):
            self._add(data)


def note_runs(note: str) -> List[Tuple[str, bool, bool, bool]]:
    parser = _Runs()
    parser.feed(note or "")
    parser.close()
    runs = parser.runs
    while runs and runs[-1][0].strip() == "":
        runs.pop()                                     # no trailing blank lines
    return runs


def _font(bold=False, italic=False, underline=False) -> InlineFont:
    return InlineFont(rFont=FONT, sz=SIZE, b=bold or None, i=italic or None, u="single" if underline else None)


def description_cell(task: Dict):
    """What goes in the description cell. A task: its description, with the factor reason under it. A note (a heading with
    no hours): its formatted text. Plain text whenever nothing in it is formatted."""
    if task.get("kind") == "note":
        runs = _no_blank_runs(note_runs(text(task.get("note"))))
        if not any(b or i or u for _, b, i, u in runs):
            return "".join(t for t, *_ in runs)
        return CellRichText(*[TextBlock(_font(b, i, u), t) for t, b, i, u in runs])
    description = text(task.get("description"))
    reason = text(task.get("factor_reason")).strip() if task.get("factor") not in (None, 1, 1.0) else ""
    if not reason:
        return description
    segments = [(description, False, False, False), ("\nFactor: " + reason, False, True, False)]
    return CellRichText(*[TextBlock(_font(b, i, u), t) for t, b, i, u in _no_blank_runs(segments)])


def _no_blank_runs(segments: List[Tuple[str, bool, bool, bool]]) -> List[Tuple[str, bool, bool, bool]]:
    """Excel refuses a rich-text run that is only white space (such as a lone line break) and calls the file corrupt.
    So white space is folded into the start of the run after it, and dropped when nothing follows."""
    out, carry = [], ""
    for value, bold, italic, underline in segments:
        if not value.strip():
            carry += value
        else:
            out.append((carry + value, bold, italic, underline))
            carry = ""
    if out:
        out[0] = (out[0][0].lstrip("\n"), *out[0][1:])
    return out


def _lines(task: Dict) -> int:
    """About how many lines the description cell needs, to set the row height (merged cells do not fit themselves)."""
    if task.get("kind") == "note":
        plain = "".join(t for t, *_ in _no_blank_runs(note_runs(text(task.get("note")))))
    else:
        plain = text(task.get("description"))
        reason = text(task.get("factor_reason")).strip() if task.get("factor") not in (None, 1, 1.0) else ""
        if reason:
            plain += "\nFactor: " + reason
    return sum(max(1, math.ceil(len(line) / DESCRIPTION_CHARS)) for line in plain.split("\n"))


# ---- the sheets ---------------------------------------------------------------------------------------
def _setup(ws) -> None:
    for col, width in enumerate(WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.4
    ws.print_options.horizontalCentered = True


def _head(ws, jc: Dict, wp: Dict, top: int) -> None:
    """The seven rows at the top of each page, starting at row `top`."""
    a, b, c, d, e, f, g = range(top, top + 7)
    put(ws, a, 1, 3, "NMC ENERGY", bold=True, size=14, align="center")
    put(ws, a, 4, 8, "CONSTRUCTION JOBCARD", bold=True, size=15, align="center")
    put(ws, a, 9, 11)
    ws.row_dimensions[a].height = 44
    put(ws, b, 1, 2, "JOBCARD NUMBER:", bold=True); put(ws, b, 3, 5, text(jc.get("number")), align="center")
    put(ws, b, 6, 6, "REVISION:", bold=True); put(ws, b, 7, 8, text(jc.get("revision")), bold=True, blue=True, align="center")
    put(ws, b, 9, 9, "FCP No.", bold=True); put(ws, b, 10, 11, text(wp.get("fcp_no")), align="center")
    put(ws, c, 1, 2, "JOBCARD TITLE:", bold=True); put(ws, c, 3, 11, text(jc.get("title")), blue=True, align="center")
    put(ws, d, 1, 3, "WORKPACK NUMBER:", bold=True); put(ws, d, 4, 11, "TITLE:", bold=True)
    put(ws, e, 1, 3, text(wp.get("number")), size=9); put(ws, e, 4, 11, text(wp.get("title")))
    ws.row_dimensions[e].height = 32
    put(ws, f, 1, 3, "ASSET:", bold=True); put(ws, f, 4, 6, "ANS SH1 NUMBER:", bold=True); put(ws, f, 7, 11, "LEAD TRADE:", bold=True)
    put(ws, g, 1, 3, text(jc.get("asset")), align="center"); put(ws, g, 4, 6, text(jc.get("ans_sh1_no")), align="center")
    put(ws, g, 7, 11, text(jc.get("trade")), blue=True, align="center")


def _details(ws, jc: Dict, wp: Dict) -> int:
    """Page 1: details, flags, drawings and completion. Returns the last row used."""
    _head(ws, jc, wp, 1)
    r = 9
    put(ws, r, 1, 11, "CREATED BY", bold=True, align="center")
    put(ws, r + 1, 1, 1, "Name"); put(ws, r + 1, 2, 5, text(jc.get("created_by_name")), blue=True, align="center")
    put(ws, r + 1, 6, 7, "Position"); put(ws, r + 1, 8, 11, text(jc.get("created_by_position")), blue=True, align="center")
    put(ws, r + 2, 1, 1, "Signature"); put(ws, r + 2, 2, 5, text(jc.get("created_by_signature")), blue=True, align="center")
    put(ws, r + 2, 6, 7, "Date"); put(ws, r + 2, 8, 11, text(jc.get("created_date")), blue=True, align="center")
    r += 4
    put(ws, r, 1, 11, "Additional information", bold=True, align="center")
    info = text(jc.get("additional_info"))
    put(ws, r + 1, 1, 11, info, blue=True, valign="top")
    ws.row_dimensions[r + 1].height = max(76, 13 * sum(max(1, math.ceil(len(x) / INFO_CHARS)) for x in info.split("\n")) + 4)
    r += 2
    spans = [(1, 2, 3), (4, 5, 6), (7, 10, 11)]                  # (label from, label to, tick column) for each group
    for i in range(7):
        for (first, last, tick), (_, flags) in zip(spans, FLAG_GROUPS):
            put(ws, r, first, last, flags[i], align="center")
            put(ws, r, tick, tick, "✓" if flags[i] in jc.get("flags", []) else None, blue=True, align="center")
        r += 1
    r += 1
    put(ws, r, 1, 11, "DRAWING REFERENCE", bold=True, align="center")
    put(ws, r + 1, 1, 3, "DRAWING NUMBER", align="center"); put(ws, r + 1, 4, 11, "TITLE", align="center")
    r += 2
    drawings = jc.get("drawings", [])
    for i in range(max(3, len(drawings))):
        d = drawings[i] if i < len(drawings) else {}
        put(ws, r, 1, 3, text(d.get("number")), align="center", size=9)
        put(ws, r, 4, 11, text(d.get("title")), blue=True, align="center", size=9)
        ws.row_dimensions[r].height = 28
        r += 1
    r += 1
    put(ws, r, 1, 11, "JOBCARD COMPLETION", bold=True, align="center")
    put(ws, r + 1, 1, 3, "COMPLETED BY", bold=True, align="center"); put(ws, r + 1, 4, 6, "SUPERVISOR", bold=True, align="center")
    put(ws, r + 1, 7, 11, "APPROVED BY (Client)", bold=True, align="center")
    r += 2
    for label in ("Name", "Signature", "Position", "Date"):
        put(ws, r, 1, 1, label); put(ws, r, 2, 3, None, blue=True)
        put(ws, r, 4, 4, label); put(ws, r, 5, 6, None, blue=True)
        put(ws, r, 7, 8, label); put(ws, r, 9, 11, None, blue=True)
        ws.row_dimensions[r].height = 20
        r += 1
    return r - 1


def _tasks(ws, jc: Dict, wp: Dict, top: int) -> Tuple[int, Dict[str, float]]:
    """Page 2: the header again, then the task table. Returns the last row used, and the result of each formula in it
    (cell reference -> number), so the results can be stored in the file beside the formulas."""
    results: Dict[str, float] = {}
    _head(ws, jc, wp, top)
    put(ws, top + 8, 1, 11, "TASKS", bold=True, align="center")
    for col, label in ((1, "ITEM"), (8, "QTY"), (9, "NORM"), (10, "FACTOR"), (11, "HOURS")):
        put(ws, top + 9, col, col, label, align="center")
    put(ws, top + 9, 2, 7, "DESCRIPTION", align="center")
    tasks = jc.get("tasks", [])
    closing = tasks[-1] if tasks and tasks[-1].get("role") == "last" else None
    body = tasks[:-1] if closing else tasks
    count = max(MIN_TASK_ROWS, len(tasks))
    first = top + 10
    for i in range(count):
        row = first + i
        t = closing if closing and i == count - 1 else (body[i] if i < len(body) else None)   # the closing task is last
        put(ws, row, 1, 1, i + 1, align="center", valign="top")
        put(ws, row, 2, 7, description_cell(t) if t else None, blue=True, valign="top")
        qty, norm = (t.get("qty"), t.get("norm_hours")) if t else (None, None)
        put(ws, row, 8, 8, qty, blue=True, align="center", valign="top")
        put(ws, row, 9, 9, norm, blue=True, align="center", valign="top")
        factor = (t.get("factor") if qty is not None or norm is not None else None) if t else None
        put(ws, row, 10, 10, factor, blue=True, align="center", valign="top")
        # quantity x norm x factor, rounded up to the next whole hour (rounded to 6 places first, as calc.line_hours does)
        put(ws, row, 11, 11, f'=IF(OR(H{row}="",I{row}=""),0,ROUNDUP(ROUND(H{row}*I{row}*IF(J{row}="",1,J{row}),6),0))',
            blue=True, align="center", valign="top", fmt=_HIDE_ZERO, formula=True)
        results[f"K{row}"] = calc.line_hours(qty, norm, factor)
        ws.row_dimensions[row].height = max(22, 13 * (_lines(t) if t else 1) + 5)
    last = first + count - 1
    put(ws, last + 1, 1, 10, "TOTAL MANHOURS:", bold=True, align="right")
    put(ws, last + 1, 11, 11, f"=SUM(K{first}:K{last})", bold=True, align="center", formula=True)
    results[f"K{last + 1}"] = sum(results.values())
    return last + 1, results


def _store_results(path: Path, results: Dict[str, float]) -> None:
    """Put each formula's result in the file beside the formula. openpyxl writes only the formula, and a viewer that does
    not calculate (Excel's Protected View, which is how a downloaded file opens, a file previewer, a PDF converter) would
    show the hours and the total empty. Excel still recalculates as soon as anything is edited."""
    sheet, temp = "xl/worksheets/sheet1.xml", path.with_name(path.name + ".tmp")
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(temp, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == sheet:
                xml = data.decode("utf-8")
                for ref, value in results.items():
                    xml = re.sub(rf'(<c r="{ref}"[^>]*><f>[^<]*</f>)(?:<v></v>|<v\s*/>)',
                                 lambda m: f"{m.group(1)}<v>{value:.10g}</v>", xml, count=1)
                data = xml.encode("utf-8")
            target.writestr(item, data)
    temp.replace(path)


def write_jobcard(path: Path, jc: Dict, wp: Dict) -> None:
    """Write one jobcard (as Repo.get_jobcard gives it) and its jobcards workpack (Repo.get_workpack) to an .xlsx file."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Jobcard"
    _setup(ws)
    page_one_end = _details(ws, jc, wp)
    ws.row_breaks.append(Break(id=page_one_end))                      # the tasks start a new page
    end, results = _tasks(ws, jc, wp, page_one_end + 1)
    ws.print_area = f"A1:K{end}"
    wb.properties.title = text(jc.get("number"))
    wb.properties.creator = "NMC Tools"
    wb.save(path)
    _store_results(path, results)
