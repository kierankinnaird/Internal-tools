import io
import sqlite3
import zipfile

import openpyxl
import pytest

import modules.workpacks.web as web
from hosts.standalone import create_app
from modules.jobcards import export
from modules.workpacks.fill import FillReport, FormResult

FOLDER = "9.0 JOBCARDS"


class StubFiller:
    """Stands in for the engine: one good form, and it remembers what it was asked to do."""
    calls = []

    def __init__(self, repo, store):
        self.repo = repo

    def fill(self, wp_id, out_dir, cert_pks=None, spare_rows=None, add_rows=False, overwrite=False):
        StubFiller.calls.append(wp_id)
        (out_dir / "AS1").mkdir()
        (out_dir / "AS1" / "AS1-1.docx").write_bytes(b"docx")
        (out_dir / "fill-log.csv").write_text("log", encoding="utf-8")
        return FillReport("Wellbay 18", out_dir, [FormResult("AS1-1", "AS1", "OK", "AS1/AS1-1.docx", "")], [])


def make_ready(tmp):
    folder = tmp / "workpacks" / "templates"
    folder.mkdir(parents=True)
    (folder / "AS1 blank.docx").write_bytes(b"x")


@pytest.fixture()
def pack(tmp_path, monkeypatch):
    monkeypatch.setattr(web, "FillService", StubFiller)
    StubFiller.calls = []
    make_ready(tmp_path)
    app = create_app(userdata=tmp_path, user_name="tester", modules=["jobcards", "workpacks"])
    c = app.test_client()
    c.post("/workpacks/", data={"title": "Wellbay 18", "jc_mode": "create", "jc_number": "BB-26-M0001", "jc_count": "2",
                                "jc_prefix": "NMC-BB18-"})
    c.post("/workpacks/1/certificates", data={"cert_id": "AS1-1", "form": "AS1"})
    return c, tmp_path, app


def fill(c, **extra):
    return c.post("/workpacks/1/fill", data=dict({"cert": ["1"]}, **extra))


def item_row(ws):
    """The row of the tasks table's column headings (the jobcard is one sheet: page 1, then the tasks)."""
    return next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "ITEM")


def actions(tmp):
    return [a for (a,) in sqlite3.connect(tmp / "audit.sqlite").execute("SELECT action FROM audit")]


def runs(tmp):
    base = tmp / "workpacks" / "output" / "wp1"
    return sorted(base.iterdir()) if base.exists() else []


# ---- the fill page ------------------------------------------------------------------------------
def test_the_fill_page_offers_the_linked_jobcards_ticked_by_default(pack):
    c, _, _ = pack
    html = c.get("/workpacks/1/fill").get_data(as_text=True)
    assert 'name="include_jobcards"' in html and 'name="include_jobcards" value="1" checked' in html
    assert "Include the linked jobcards (2 in 1 jobcards workpack)" in html and "9.0 JOBCARDS" in html


def test_the_page_counts_across_several_jobcards_workpacks(pack):
    c, _, _ = pack
    c.post("/workpacks/1/jobcards/create", data={"jc_number": "SECOND", "jc_count": "1", "jc_prefix": "NMC-BB19-"})
    assert "Include the linked jobcards (3 in 2 jobcards workpacks)" in c.get("/workpacks/1/fill").get_data(as_text=True)


def test_with_nothing_linked_the_page_says_so_and_offers_no_tick_box(pack):
    c, _, _ = pack
    c.post("/workpacks/", data={"title": "Second"})
    c.post("/workpacks/2/certificates", data={"cert_id": "AS1-1", "form": "AS1"})
    html = c.get("/workpacks/2/fill").get_data(as_text=True)
    assert 'name="include_jobcards"' not in html and "No jobcards are linked to this workpack" in html


def test_without_the_jobcards_module_there_is_no_option_and_filling_is_unchanged(tmp_path, monkeypatch):
    monkeypatch.setattr(web, "FillService", StubFiller)
    make_ready(tmp_path)
    c = create_app(userdata=tmp_path, user_name="tester", modules=["workpacks"]).test_client()
    c.post("/workpacks/", data={"title": "W"})
    c.post("/workpacks/1/certificates", data={"cert_id": "AS1-1", "form": "AS1"})
    html = c.get("/workpacks/1/fill").get_data(as_text=True)
    assert "include_jobcards" not in html and "Jobcards" not in html and "9.0" not in html
    r = c.post("/workpacks/1/fill", data={"cert": ["1"], "include_jobcards": "1"})          # a stray field is ignored
    assert r.status_code == 302 and "error" not in r.headers["Location"]
    assert not any(p.name == FOLDER for p in (tmp_path / "workpacks" / "output" / "wp1").rglob("*"))


def test_without_jobcards_view_the_option_is_hidden(pack):
    c, _, app = pack
    app.config["CTX"]._permissions = {"workpacks.view", "workpacks.fill"}
    assert "include_jobcards" not in c.get("/workpacks/1/fill").get_data(as_text=True)


# ---- filling ------------------------------------------------------------------------------------
def test_the_jobcards_are_written_into_the_run_as_excel_files(pack):
    c, tmp, _ = pack
    r = fill(c, include_jobcards="1")
    assert r.status_code == 302 and StubFiller.calls == [1]
    (run,) = runs(tmp)
    assert sorted(p.name for p in (run / FOLDER).iterdir()) == ["NMC-BB18-001.xlsx", "NMC-BB18-002.xlsx"]
    assert (run / "AS1" / "AS1-1.docx").exists()                                         # the forms are unaffected
    wb = openpyxl.load_workbook(run / FOLDER / "NMC-BB18-002.xlsx")
    assert wb["Jobcard"]["C2"].value == "NMC-BB18-002" and wb["Jobcard"]["A5"].value == "BB-26-M0001"
    assert wb.sheetnames == ["Jobcard"]                                                  # one sheet, so Print gets every page
    assert wb["Jobcard"].cell(item_row(wb["Jobcard"]) + 1, 2).value == "Raise PTW, Carry out TBT, Erect Signs & Barriers"
    assert "jobcards_added_to_run" in actions(tmp)


def test_the_result_page_lists_each_jobcard_with_a_download_link(pack):
    c, _, _ = pack
    html = c.get(fill(c, include_jobcards="1").headers["Location"]).get_data(as_text=True)
    assert "3 of 3 forms written" in html
    assert html.count("<td>Jobcard</td>") == 2 and "NMC-BB18-001" in html and f"{FOLDER}/NMC-BB18-001.xlsx" in html
    assert "AS1/AS1-1.docx" in html


def test_a_jobcard_file_can_be_downloaded_on_its_own(pack):
    c, _, _ = pack
    loc = fill(c, include_jobcards="1").headers["Location"]
    r = c.get(f"{loc}/files/{FOLDER}/NMC-BB18-001.xlsx")
    assert r.status_code == 200 and "attachment" in r.headers["Content-Disposition"]
    assert openpyxl.load_workbook(io.BytesIO(r.data))["Jobcard"]["C2"].value == "NMC-BB18-001"


def test_the_zip_holds_the_jobcards_in_their_own_folder_beside_the_forms(pack):
    c, _, _ = pack
    loc = fill(c, include_jobcards="1").headers["Location"]
    names = zipfile.ZipFile(io.BytesIO(c.get(f"{loc}/download").data)).namelist()
    assert f"{FOLDER}/NMC-BB18-001.xlsx" in names and f"{FOLDER}/NMC-BB18-002.xlsx" in names
    assert "AS1/AS1-1.docx" in names and "fill-log.csv" in names and "run.json" not in names


def test_unticked_means_no_jobcards(pack):
    c, tmp, _ = pack
    loc = fill(c).headers["Location"]                                                    # the box was not ticked
    (run,) = runs(tmp)
    assert not (run / FOLDER).exists()
    html = c.get(loc).get_data(as_text=True)
    assert "<td>Jobcard</td>" not in html and "1 of 1 form written" in html
    assert "jobcards_added_to_run" not in actions(tmp)


def test_only_jobcards_linked_to_this_workpack_are_included(pack):
    c, tmp, _ = pack
    c.post("/workpacks/", data={"title": "Other", "jc_mode": "create", "jc_number": "OTHER", "jc_count": "3", "jc_prefix": "X-"})
    fill(c, include_jobcards="1")
    (run,) = runs(tmp)
    assert sorted(p.name for p in (run / FOLDER).iterdir()) == ["NMC-BB18-001.xlsx", "NMC-BB18-002.xlsx"]


def test_jobcards_are_written_as_they_are_saved_at_that_moment(pack):
    c, tmp, app = pack
    c.post("/jobcards/cards/1/save", json={"title": "Edited just now", "tasks": [
        {"kind": "note", "note": "<b>Destruct tubing to:</b>"},
        {"description": "Annulus Kill MSAS Valve", "qty": 2, "norm_hours": 3, "basis": "q"}]})
    fill(c, include_jobcards="1")
    (run,) = runs(tmp)
    ws = openpyxl.load_workbook(run / FOLDER / "NMC-BB18-001.xlsx")["Jobcard"]
    assert ws["C3"].value == "Edited just now"
    head = item_row(ws)
    assert str(ws.cell(head + 2, 2).value) == "Destruct tubing to:" and ws.cell(head + 2, 8).value is None       # the note: no hours
    assert ws.cell(head + 3, 2).value == "Annulus Kill MSAS Valve" and ws.cell(head + 3, 8).value == 2             # its subtask has them


def test_one_jobcard_failing_is_reported_and_the_rest_are_still_there(pack, monkeypatch):
    c, tmp, _ = pack
    real = export.write_jobcard

    def flaky(path, jc, wp):
        if jc["number"] == "NMC-BB18-001":
            raise RuntimeError("cannot write")
        return real(path, jc, wp)
    monkeypatch.setattr(export, "write_jobcard", flaky)
    html = c.get(fill(c, include_jobcards="1").headers["Location"]).get_data(as_text=True)
    assert "2 of 3 forms written, 1 need attention" in html and "cannot write" in html
    (run,) = runs(tmp)
    assert sorted(p.name for p in (run / FOLDER).iterdir()) == ["NMC-BB18-002.xlsx"]


def test_if_the_jobcards_cannot_be_added_the_forms_are_still_kept(pack, monkeypatch):
    c, tmp, app = pack
    service = app.config["CTX"].services.get("jobcards.link")

    def boom(*a, **k):
        raise RuntimeError("secret internal detail")
    monkeypatch.setattr(service, "export_linked", boom)
    r = fill(c, include_jobcards="1")
    html = c.get(r.headers["Location"]).get_data(as_text=True)
    assert "The jobcards could not be written, so they are not in this run" in html and "secret internal detail" not in html
    assert "1 of 1 form written" in html
    assert (runs(tmp)[0] / "AS1" / "AS1-1.docx").exists()


def test_asking_for_jobcards_without_permission_is_refused_before_anything_is_written(pack):
    c, tmp, app = pack
    app.config["CTX"]._permissions = {"workpacks.view", "workpacks.fill"}
    assert fill(c, include_jobcards="1").status_code == 403
    assert StubFiller.calls == [] and runs(tmp) == []
    assert fill(c).status_code == 302                                                    # filling without them is still fine


def test_a_failed_fill_does_not_leave_jobcards_behind(pack, monkeypatch):
    c, tmp, app = pack

    class Refusing(StubFiller):
        def fill(self, *a, **k):
            raise ValueError("The blank templates are not installed")
    monkeypatch.setattr(web, "FillService", Refusing)
    c2 = create_app(userdata=tmp, user_name="tester", modules=["jobcards", "workpacks"]).test_client()
    r = c2.post("/workpacks/1/fill", data={"cert": ["1"], "include_jobcards": "1"})
    assert "error=" in r.headers["Location"] and runs(tmp) == []


def test_titles_with_markup_do_not_break_the_result_page(pack):
    c, _, _ = pack
    c.post("/jobcards/cards/1/save", json={"title": "<script>alert(1)</script>", "tasks": []})
    html = c.get(fill(c, include_jobcards="1").headers["Location"]).get_data(as_text=True)
    assert "<script>alert(1)</script>" not in html
