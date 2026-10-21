"""Estimates (a workbook of CTRs): storage, and the workbook file with its Summary sheet and one sheet per CTR."""
import re
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

import openpyxl
import pytest

from core.migrations import migrate
from hosts.standalone import create_app
from modules.ctrs import calc, export, repo as repo_module
from modules.ctrs.constants import DISCIPLINES
from modules.ctrs.repo import Repo

MIGRATIONS = Path(__file__).resolve().parent.parent / "modules" / "ctrs" / "migrations"


@pytest.fixture
def repo(tmp_path):
    migrate(tmp_path / "ctrs.sqlite", "ctrs", MIGRATIONS)
    log = []
    r = Repo(tmp_path / "ctrs.sqlite", audit=lambda a, d: log.append((a, d)))
    r.log = log
    return r


def piping(**kw):
    body = {"title": "Piping", "lines": [{"description": "P/Fitter", "grp": "P/Fitter", "eng_override": "137"},
                                         {"description": "Isometrics", "flag": "D", "qty": "17", "eng_norm": "3"}],
            "resources": [{"name": "Lead", "source": "group:P/Fitter", "split_pct": "60", "rate": "55.5"},
                          {"name": "Senior", "source": "group:P/Fitter", "split_pct": "40", "rate": "48"},
                          {"name": "Designer", "source": "pool1", "split_pct": "100", "rate": "60"}],
            "additions": [{"name": "Offshore", "unit_cost": "100", "units": "3"}]}
    body.update(kw)
    return body


# ---- storage ------------------------------------------------------------------------------------------------
def test_a_workbook_can_start_with_a_ctr_for_each_discipline_in_the_usual_order(repo):
    eid = repo.create_estimate({"title": "BB diesel", "client": "Apache"}, "kieran", ["Piping", "Process", "Mechanical"])
    e = repo.get_estimate(eid)
    assert [c["title"] for c in e["ctrs"]] == ["Process", "Piping", "Mechanical"]
    assert {c["client"] for c in e["ctrs"]} == {"Apache"} and {c["estimate_id"] for c in e["ctrs"]} == {eid}
    assert repo.log[0] == ("estimate_created", "BB diesel") and [a for a, _ in repo.log].count("ctr_created") == 3


def test_every_discipline_on_offer_can_be_chosen_at_once(repo):
    eid = repo.create_estimate({"title": "All"}, "k", DISCIPLINES)
    assert [c["title"] for c in repo.get_estimate(eid)["ctrs"]] == DISCIPLINES


def test_an_unknown_discipline_or_a_missing_name_makes_nothing(repo):
    with pytest.raises(ValueError, match="Unknown discipline: Basket weaving"):
        repo.create_estimate({"title": "BB"}, "k", ["Piping", "Basket weaving"])
    with pytest.raises(ValueError, match="Estimate name is required"):
        repo.create_estimate({"title": "  "}, "k")
    with pytest.raises(ValueError, match="too long"):
        repo.create_estimate({"title": "x" * 300}, "k")
    assert repo.list_estimates() == [] and repo.list_ctrs() == [] and repo.log == []


def test_a_ctr_takes_the_project_details_of_its_workbook(repo):
    eid = repo.create_estimate({"title": "BB", "client": "Apache", "project_name": "Beryl", "framework_ref": "FA-9"}, "k")
    cid = repo.create_ctr({"title": "Piping", "client": "Somebody else"}, "k", estimate_id=eid)
    ctr = repo.get_ctr(cid)
    assert (ctr["client"], ctr["project_name"], ctr["framework_ref"]) == ("Apache", "Beryl", "FA-9")
    with pytest.raises(ValueError, match="No such estimate"):
        repo.create_ctr({"title": "x"}, "k", estimate_id=99)


def test_a_ctr_made_on_its_own_gets_a_workbook_named_after_it_without_noise_in_the_log(repo):
    cid = repo.create_ctr({"title": "Piping", "client": "Apache"}, "k")
    e = repo.get_estimate(repo.get_ctr(cid)["estimate_id"])
    assert e["title"] == "Piping" and e["client"] == "Apache" and [c["id"] for c in e["ctrs"]] == [cid]
    assert repo.log == [("ctr_created", "Piping")]


def test_saving_the_workbook_copies_the_project_details_to_every_sheet(repo):
    eid = repo.create_estimate({"title": "BB"}, "k", ["Piping", "Process"])
    repo.save_estimate(eid, {"title": "BB diesel", "client": "Apache", "project_name": "Beryl", "project_title": "Diesel",
                             "framework_ref": "FA-9"}, "k")
    for ctr in repo.get_estimate(eid)["ctrs"]:
        assert (ctr["client"], ctr["project_name"], ctr["project_title"], ctr["framework_ref"]) == ("Apache", "Beryl", "Diesel", "FA-9")
    assert repo.get_estimate(eid)["title"] == "BB diesel" and ("estimate_saved", "BB diesel") in repo.log
    other = repo.create_estimate({"title": "Other", "client": "Shell"}, "k", ["Piping"])
    repo.save_estimate(eid, {"title": "BB diesel", "client": "Apache 2"}, "k")
    assert repo.get_estimate(other)["ctrs"][0]["client"] == "Shell"                      # another workbook is untouched


def test_a_refused_workbook_save_changes_nothing(repo):
    eid = repo.create_estimate({"title": "BB", "client": "Apache"}, "k", ["Piping"])
    with pytest.raises(ValueError, match="Estimate name is required"):
        repo.save_estimate(eid, {"title": " ", "client": "Changed"}, "k")
    with pytest.raises(ValueError, match="No such estimate"):
        repo.save_estimate(99, {"title": "x"}, "k")
    assert repo.get_estimate(eid)["client"] == "Apache" and repo.get_estimate(eid)["ctrs"][0]["client"] == "Apache"


def test_the_workbook_totals_add_up_its_ctrs(repo):
    eid = repo.create_estimate({"title": "BB"}, "k")
    a, b = (repo.create_ctr({"title": t}, "k", estimate_id=eid) for t in ("Piping", "Mechanical"))
    repo.save_ctr(a, piping(), "k")
    repo.save_ctr(b, piping(title="Mechanical"), "k")
    one = calc.compute(repo.get_ctr(a))["totals"]
    e = repo.get_estimate(eid)
    assert e["hours"] == 2 * one["hours"] and e["cost"] == round(2 * one["cost"], 2)
    assert [x["id"] for x in repo.list_estimates()] == [eid]


def test_a_workbook_holds_a_limited_number_of_ctrs(repo, monkeypatch):
    monkeypatch.setattr(repo_module, "MAX_CTRS", 2)
    eid = repo.create_estimate({"title": "BB"}, "k", ["Piping", "Process"])
    with pytest.raises(ValueError, match="at most 2 CTRs"):
        repo.create_ctr({"title": "Third"}, "k", estimate_id=eid)


def test_deleting_a_workbook_removes_its_ctrs_and_leaves_others(repo):
    gone = repo.create_estimate({"title": "Gone"}, "k", ["Piping"])
    kept = repo.create_estimate({"title": "Kept"}, "k", ["Piping"])
    repo.delete_estimate(gone)
    assert repo.get_estimate(gone) is None and [e["id"] for e in repo.list_estimates()] == [kept]
    assert len(repo.list_ctrs()) == 1 and ("estimate_deleted", "Gone") in repo.log


def test_a_duplicate_stays_in_the_same_workbook(repo):
    eid = repo.create_estimate({"title": "BB", "client": "Apache"}, "k", ["Piping"])
    copy = repo.duplicate_ctr(repo.get_estimate(eid)["ctrs"][0]["id"], "k")
    assert repo.get_ctr(copy)["estimate_id"] == eid and [c["title"] for c in repo.get_estimate(eid)["ctrs"]] == ["Piping", "Piping (copy)"]


# ---- the workbook file --------------------------------------------------------------------------------------
@pytest.fixture
def built(repo, tmp_path):
    eid = repo.create_estimate({"title": "BB diesel", "client": "Apache", "project_name": "Beryl Bravo",
                                "project_title": "Diesel mods", "framework_ref": "FA-9"}, "k")
    for title in ("Piping", "Piping", "Mechanical"):                                        # two with the same name
        cid = repo.create_ctr({"title": title, "ctr_no": "C-" + title[:2]}, "k", estimate_id=eid)
        repo.save_ctr(cid, piping(title=title), "k")
    estimate = repo.get_estimate(eid)
    ctrs = [repo.get_ctr(c["id"]) for c in estimate["ctrs"]]
    path = tmp_path / "wb.xlsx"
    export.write_workbook(path, estimate, ctrs)
    return estimate, ctrs, path


def cached_in(path, n):
    xml = zipfile.ZipFile(path).read(f"xl/worksheets/sheet{n}.xml").decode("utf-8")
    return {m.group(1): float(m.group(2)) for m in re.finditer(r'<c r="([A-Z]+\d+)"[^>]*><f>[^<]*</f><v>([^<]*)</v>', xml)}


def test_the_workbook_has_a_summary_then_a_sheet_for_each_ctr_with_unique_names(built):
    _, _, path = built
    assert openpyxl.load_workbook(path).sheetnames == ["Summary", "Piping", "Piping (2)", "Mechanical"]


@pytest.mark.parametrize("titles, expected", [
    (["Summary"], ["Summary", "Summary (2)"]),                                              # a CTR may not take the summary's name
    (["piping", "Piping", "PIPING"], ["Summary", "piping", "Piping (2)", "PIPING (3)"]),    # Excel ignores case
    (["x" * 40, "x" * 40], ["Summary", "x" * 31, "x" * 25 + " (2)" if False else "x" * 27 + " (2)"]),
    (["It's / a: test?"], ["Summary", "It s   a  test"]),                                    # no apostrophe or other banned character
])
def test_sheet_names_are_always_acceptable_to_excel(titles, expected):
    assert ["Summary"] + export.unique_names(titles, taken=("Summary",)) == expected


def test_the_summary_shows_the_project_details_and_a_line_for_each_ctr(built):
    estimate, ctrs, path = built
    ws = openpyxl.load_workbook(path)["Summary"]
    assert ws["A1"].value == "CTR Summary" and ws["A2"].value == "BB diesel"
    shown = {ws.cell(r, 1).value: ws.cell(r, 3).value for r in range(3, 7)}
    assert shown == {"Client": "Apache", "Project Name": "Beryl Bravo", "Project Title": "Diesel mods", "Framework Ref": "FA-9"}
    rows = [(ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 3).value) for r in range(9, 12)]
    assert rows == [(1, "Piping", "C-Pi"), (2, "Piping", "C-Pi"), (3, "Mechanical", "C-Me")]


def test_the_summary_figures_are_formulas_that_point_at_each_sheets_own_totals(built):
    _, ctrs, path = built
    wb = openpyxl.load_workbook(path)
    ws = wb["Summary"]
    for i, (name, ctr) in enumerate(zip(wb.sheetnames[1:], ctrs)):
        r = 9 + i
        hours = re.fullmatch(rf"='{re.escape(name)}'!(E\d+)", ws.cell(r, 4).value)
        cost = re.fullmatch(rf"='{re.escape(name)}'!(I\d+)", ws.cell(r, 5).value)
        assert hours and cost, (name, ws.cell(r, 4).value, ws.cell(r, 5).value)
        assert wb[name][hours.group(1)].value.startswith("=SUM(E") and wb[name][cost.group(1)].value.startswith("=SUM(I")
        assert ws.cell(r, 2).hyperlink.location == f"'{name}'!A1"                            # the name jumps to the sheet
    assert ws["D12"].value == "=SUM(D9:D11)" and ws["E12"].value == "=SUM(E9:E11)" and ws["A12"].value == "Total"


def test_every_stored_result_is_there_so_protected_view_shows_numbers(built):
    _, ctrs, path = built
    totals = [calc.compute(c)["totals"] for c in ctrs]
    summary = cached_in(path, 1)
    assert [summary[f"D{r}"] for r in (9, 10, 11)] == [t["hours"] for t in totals]
    assert [summary[f"E{r}"] for r in (9, 10, 11)] == pytest.approx([t["cost"] for t in totals])
    assert summary["D12"] == sum(t["hours"] for t in totals) and summary["E12"] == pytest.approx(sum(t["cost"] for t in totals))
    for n, t in enumerate(totals, 2):                                                       # and on each CTR sheet
        results = cached_in(path, n)
        assert t["hours"] in results.values() and any(abs(v - t["cost"]) < 0.005 for v in results.values())
    values = openpyxl.load_workbook(path, data_only=True)["Summary"]                        # a reader that does not calculate
    assert values["D12"].value == sum(t["hours"] for t in totals)


def test_each_ctr_sheet_has_the_logo_and_the_summary_does_not(built):
    _, _, path = built
    z = zipfile.ZipFile(path)
    names = z.namelist()
    assert z.testzip() is None and len(names) == len(set(names))
    for n in (2, 3, 4):
        assert f"xl/drawings/drawing{n}.xml" in names and f"xl/drawings/_rels/drawing{n}.xml.rels" in names
        assert f"drawing{n}.xml" in z.read(f"xl/worksheets/_rels/sheet{n}.xml.rels").decode()
        assert z.read(f"xl/worksheets/sheet{n}.xml").decode().count("<drawing ") == 1
        assert f"/xl/drawings/drawing{n}.xml".encode() in z.read("[Content_Types].xml")
    assert "xl/drawings/drawing1.xml" not in names and b"<drawing" not in z.read("xl/worksheets/sheet1.xml")
    assert names.count("xl/media/image1.png") == 1 and z.read("xl/media/image1.png") == export.LOGO.read_bytes()
    for name in names:
        if name.endswith((".xml", ".rels")):
            ET.fromstring(z.read(name))                                                     # every part is well-formed
    assert b'TargetMode="External"' not in b"".join(z.read(n) for n in names if n.endswith(".rels"))   # no link leaves the file


def test_every_ctr_sheet_shows_the_workbooks_project_details(built):
    _, _, path = built
    wb = openpyxl.load_workbook(path)
    for name in wb.sheetnames[1:]:
        ws = wb[name]
        assert (ws["B5"].value, ws["B6"].value, ws["B7"].value) == ("Beryl Bravo", "Diesel mods", "Apache")


def test_an_empty_workbook_still_makes_a_good_summary(repo, tmp_path):
    estimate = repo.get_estimate(repo.create_estimate({"title": "Empty"}, "k"))
    export.write_workbook(tmp_path / "e.xlsx", estimate, [])
    wb = openpyxl.load_workbook(tmp_path / "e.xlsx")
    assert wb.sheetnames == ["Summary"] and wb["Summary"]["D10"].value == "=SUM(D9:D9)"
    assert cached_in(tmp_path / "e.xlsx", 1)["D10"] == 0


def test_a_missing_logo_gives_a_good_workbook_without_pictures(built, monkeypatch, tmp_path):
    estimate, ctrs, _ = built
    monkeypatch.setattr(export, "LOGO", tmp_path / "nope.png")
    export.write_workbook(tmp_path / "nologo.xlsx", estimate, ctrs)
    z = zipfile.ZipFile(tmp_path / "nologo.xlsx")
    assert not [n for n in z.namelist() if "drawing" in n or "media" in n] and cached_in(tmp_path / "nologo.xlsx", 2)


def test_text_that_looks_like_a_formula_stays_text_on_the_summary(repo, tmp_path):
    eid = repo.create_estimate({"title": "=1+1", "client": "=HYPERLINK(\"x\")"}, "k", [])
    cid = repo.create_ctr({"title": "=SUM(1)", "ctr_no": "=2"}, "k", estimate_id=eid)
    export.write_workbook(tmp_path / "t.xlsx", repo.get_estimate(eid), [repo.get_ctr(cid)])
    ws = openpyxl.load_workbook(tmp_path / "t.xlsx")["Summary"]
    assert ws["A2"].data_type == "s" and ws["C3"].data_type == "s" and ws["B9"].data_type == "s" and ws["C9"].data_type == "s"


# ---- the download -------------------------------------------------------------------------------------------
@pytest.fixture
def client(tmp_path):
    app = create_app(tmp_path, "tester", modules=["ctrs"])
    c = app.test_client()
    c.post("/ctrs/", data={"title": "BB / diesel: mods", "client": "Apache", "disciplines": ["Piping", "Mechanical"]})
    return c


def test_the_workbook_downloads_as_one_xlsx_named_after_the_workbook(client, tmp_path):
    r = client.get("/ctrs/estimates/1/export.xlsx")
    assert r.status_code == 200 and "spreadsheetml" in r.headers["Content-Type"]
    assert 'filename="BB diesel mods.xlsx"' in r.headers["Content-Disposition"]
    import io
    assert openpyxl.load_workbook(io.BytesIO(r.data)).sheetnames == ["Summary", "Piping", "Mechanical"]


def test_the_download_is_made_from_what_is_saved_and_leaves_no_temporary_files(client, tmp_path):
    client.post("/ctrs/1/save", json=piping())
    import io
    ws = openpyxl.load_workbook(io.BytesIO(client.get("/ctrs/estimates/1/export.xlsx").data), data_only=True)["Summary"]
    assert ws["D9"].value > 0 and ws["E9"].value > 0 and ws["D10"].value == 0          # the saved CTR has hours, the untouched one none
    assert ws["D11"].value == ws["D9"].value
    assert not list(Path(tmp_path).rglob("*.xlsx*")) and not list(Path(tmp_path).rglob("*.tmp"))


def test_a_workbook_with_no_ctrs_has_no_download_and_says_why(client):
    client.post("/ctrs/", data={"title": "Empty"})
    r = client.get("/ctrs/estimates/2/export.xlsx")
    assert r.status_code == 302 and b"Add a CTR before downloading" in client.get(r.headers["Location"]).data
    assert client.get("/ctrs/estimates/99/export.xlsx").status_code == 404


def test_a_viewer_can_download_but_not_change_a_workbook(tmp_path):
    app = create_app(tmp_path, "tester", modules=["ctrs"])
    c = app.test_client()
    c.post("/ctrs/", data={"title": "BB", "disciplines": ["Piping"]})
    app.config["CTX"]._permissions = {"ctrs.view"}
    assert c.get("/ctrs/estimates/1").status_code == 200 and c.get("/ctrs/estimates/1/export.xlsx").status_code == 200
    page = c.get("/ctrs/estimates/1").data.decode()
    assert "Add sheet" not in page and "Delete this workbook" not in page and "Save project details" not in page
    for url, data in (("/ctrs/", {"title": "x"}), ("/ctrs/estimates/1", {"title": "x"}), ("/ctrs/estimates/1/ctrs", {"title": "x"}),
                      ("/ctrs/estimates/1/delete", {})):
        assert c.post(url, data=data).status_code == 403, url
    app.config["CTX"]._permissions = set()
    assert c.get("/ctrs/estimates/1").status_code == 403 and c.get("/ctrs/estimates/1/export.xlsx").status_code == 403


def test_whatever_is_typed_into_a_workbook_cannot_break_out_of_its_pages(tmp_path):
    app = create_app(tmp_path, "tester", modules=["ctrs"])
    c = app.test_client()
    evil = "</script><script>alert(1)</script><img src=x onerror=alert(2)>"
    c.post("/ctrs/", data={"title": evil, "client": evil, "disciplines": ["Piping"]})
    for url in ("/ctrs/", "/ctrs/estimates/1", "/ctrs/1"):
        page = c.get(url).data.decode()
        assert "<script>alert" not in page and "<img src=x" not in page, url
