import io
import sqlite3
import zipfile
from pathlib import Path

import openpyxl
import pytest

from core.migrations import migrate
from hosts.standalone import create_app
from modules.workpacks import importer
from modules.workpacks.engine import fill_workpack as F
from modules.workpacks.fields import CERTIFICATE_COLUMNS, WORKPACK_COLUMNS
from modules.workpacks.importer import ImportProblem, import_workbook, parse_workbook
from modules.workpacks.repo import Repo

MIGRATIONS = Path(__file__).resolve().parent.parent / "modules" / "workpacks" / "migrations"


def new_repo(tmp_path, name="w.sqlite"):
    db = tmp_path / name
    migrate(db, "workpacks", MIGRATIONS)
    events = []
    r = Repo(db, audit=lambda a, d: events.append((a, d)))
    r.events = events
    return r


def sample(repo):
    """A workpack with every kind of table, built the way the screens build one."""
    wp = repo.create_workpack({"title": "Wellbay 18 clearance", "project_no": "FCP-123", "workpack_no": "WPK-018",
                               "site": "Alpha", "handover_includes": "Line one\n\nLine two", "revision": "B"}, "k")
    repo.add_certificate(wp, {"cert_id": "AS1-1", "form": "AS1", "equipment": "Pump P-1", "scope": "Check =1+1\n\nalignment",
                              "folder": "Mech", "filename": "AS1 pump"}, "k")
    sl1 = repo.add_certificate(wp, {"cert_id": "SL1-1", "form": "SL1", "sl1a": "SL1A-AB-12345"}, "k")
    ji1 = repo.add_certificate(wp, {"cert_id": "JI1-1", "form": "JI1"}, "k")
    ee12 = repo.add_certificate(wp, {"cert_id": "EE12-1", "form": "EE12"}, "k")
    repo.replace_rows(sl1, "line", [{"pid": "P-1", "line_no": "PL-33000-6\u201d-AHP95-G-22/ PG-32010-6\u201d-AHP95-G-22",
                                     "from_equipment": "A", "to_equipment": "B"}, {"line_no": "=SUM(1)"}], "k")
    repo.replace_rows(ji1, "joint", [{"joint_no": "J-1", "drawing": "D1", "size": "6", "rating": "300#",
                                      "joint_type": "RTJ", "torque": "230"}], "k")
    repo.replace_rows(ee12, "tag", [{"tag": "TAG-1", "drawing": "D9"}, {"tag": "TAG-2"}], "k")
    return wp


def workbook_bytes(tmp_path, repo, wp):
    """The engine's input workbook for a saved workpack, as bytes (a real round trip)."""
    from modules.workpacks.fill import FillService
    from modules.workpacks.store import TemplateStore
    path = tmp_path / "in.xlsx"
    FillService(repo, TemplateStore(tmp_path / "none")).build_workbook(wp, path)
    return path.read_bytes()


def snapshot(repo, wp):
    w = repo.get_workpack(wp)
    certs = []
    for c in repo.list_certificates(wp):
        rows = {k: repo.get_rows(c["id"], k) for k in ("line", "joint", "tag")}
        certs.append(({k: c[k] for k in ["cert_id", "form"] + CERTIFICATE_COLUMNS}, rows))
    return {k: w[k] for k in WORKPACK_COLUMNS}, certs


def strip_ids(snap):
    """Rows carry database ids that differ between two databases; drop them."""
    work, certs = snap
    clean = []
    for fields, rows in certs:
        clean.append((fields, {k: [{c: v for c, v in r.items() if c not in ("id", "certificate_id")} for r in rs]
                               for k, rs in rows.items()}))
    return work, clean


# ---- the round trip ------------------------------------------------------------------------------
def test_a_workpack_survives_going_through_a_workbook_and_back(tmp_path):
    a = new_repo(tmp_path, "a.sqlite")
    wp = sample(a)
    data = workbook_bytes(tmp_path, a, wp)
    b = new_repo(tmp_path, "b.sqlite")
    result = import_workbook(b, data, "round trip.xlsx", "kieran")
    assert (result.certificates, result.rows) == (4, 5)
    assert strip_ids(snapshot(b, result.wp_id)) == strip_ids(snapshot(a, wp))
    assert ("workpack_imported", "Wellbay 18 clearance: 4 certificates, 5 table rows, from round trip.xlsx") in b.events


def test_blank_lines_and_formula_like_text_come_through_as_typed(tmp_path):
    a = new_repo(tmp_path, "a.sqlite")
    wp = sample(a)
    b = new_repo(tmp_path, "b.sqlite")
    got = import_workbook(b, workbook_bytes(tmp_path, a, wp), "x.xlsx", "k").wp_id
    assert b.get_workpack(got)["handover_includes"] == "Line one\n\nLine two"
    assert b.list_certificates(got)[0]["scope"] == "Check =1+1\n\nalignment"
    lines = b.get_rows(b.list_certificates(got)[1]["id"], "line")
    assert lines[1]["line_no"] == "=SUM(1)"


def test_the_people_stamp_is_the_importing_user(tmp_path):
    a = new_repo(tmp_path, "a.sqlite")
    data = workbook_bytes(tmp_path, a, sample(a))
    b = new_repo(tmp_path, "b.sqlite")
    got = import_workbook(b, data, "x.xlsx", "kieran").wp_id
    assert b.get_workpack(got)["created_by"] == "kieran"


# ---- refusals ------------------------------------------------------------------------------------
def make_book(tmp_path, mutate=None):
    path = tmp_path / "book.xlsx"
    F.make_sheet(str(path))
    wb = openpyxl.load_workbook(path)
    wb["Project"].cell(2, 2).value = "Title"
    if mutate:
        mutate(wb)
    wb.save(path)
    return path.read_bytes()


def certs_sheet(wb, *rows):
    ws = wb["Certificates"]
    for i, (cid, form) in enumerate(rows, 3):
        ws.cell(i, 1).value, ws.cell(i, 2).value = cid, form


def test_a_minimal_workbook_imports(tmp_path):
    r = new_repo(tmp_path)
    got = import_workbook(r, make_book(tmp_path, lambda wb: certs_sheet(wb, ("AS1-1", "as1"))), "m.xlsx", "k")
    assert (got.certificates, got.rows) == (1, 0)
    assert r.list_certificates(got.wp_id)[0]["form"] == "AS1"           # form codes are upper-cased


def problems_of(tmp_path, data, name="b.xlsx"):
    with pytest.raises(ImportProblem) as e:
        parse_workbook(data, name)
    return e.value.problems


def test_the_engines_own_checks_are_reported_one_by_one(tmp_path):
    def mutate(wb):
        certs_sheet(wb, ("A", "AS1"), ("A", "AS1"), ("B", ""))
        wb["Lines"].cell(3, 1).value = "NOPE"
    found = problems_of(tmp_path, make_book(tmp_path, mutate))
    text = " | ".join(found)
    assert "duplicate ID 'A'" in text and "Form is blank" in text and "'NOPE' is not on the Certificates sheet" in text
    assert not any(p.startswith("-") or "problems in the input workbook" in p for p in found)


def test_a_missing_title_and_unsafe_names_are_refused(tmp_path):
    def mutate(wb):
        wb["Project"].cell(2, 2).value = None
        certs_sheet(wb, ("A", "AS1"))
        headers = {c.value: c.column for c in wb["Certificates"][1]}
        wb["Certificates"].cell(3, headers["Output folder (optional)"]).value = "../elsewhere"
        wb["Certificates"].cell(3, headers["Output file name (optional)"]).value = "C:\\x"
    found = problems_of(tmp_path, make_book(tmp_path, mutate))
    assert any("no Project title" in p for p in found)
    assert sum("plain name" in p for p in found) == 2


def test_limits_on_certificates_and_rows(tmp_path, monkeypatch):
    monkeypatch.setattr(importer, "MAX_CERTIFICATES", 2)
    found = problems_of(tmp_path, make_book(tmp_path, lambda wb: certs_sheet(wb, ("A", "AS1"), ("B", "AS1"), ("C", "AS1"))))
    assert any("limit is 2" in p for p in found)
    monkeypatch.setattr(importer, "MAX_ROWS", 1)

    def mutate(wb):
        certs_sheet(wb, ("S", "SL1"))
        wb["Lines"].cell(3, 1).value = "S"
        wb["Lines"].cell(4, 1).value = "S"
    assert any("limit is 1" in p for p in problems_of(tmp_path, make_book(tmp_path, mutate)))


def test_not_a_workbook(tmp_path):
    for data, name, expect in [(b"", "a.xlsx", "empty"), (b"hello", "a.xlsx", "not a readable"),
                               (b"hello", "a.csv", "saved as .xlsx"), (b"hello", "a.xlsm", "saved as .xlsx")]:
        assert expect in problems_of(tmp_path, data, name)[0]
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("hello.txt", "hi")
    assert "not a readable" in problems_of(tmp_path, buf.getvalue())[0]


def test_oversized_and_zip_bomb_uploads_are_not_opened(tmp_path, monkeypatch):
    data = make_book(tmp_path)
    monkeypatch.setattr(importer, "MAX_BYTES", 100)
    assert "larger than 0 KB" in problems_of(tmp_path, data)[0]
    monkeypatch.undo()
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("xl/workbook.xml", "<x/>")
        z.writestr("xl/big.bin", b"\0" * (importer.MAX_UNPACKED + 1))
    assert len(buf.getvalue()) < importer.MAX_BYTES
    assert "unreasonable size" in problems_of(tmp_path, buf.getvalue())[0]


def test_a_failure_part_way_leaves_nothing_behind(tmp_path, monkeypatch):
    a = new_repo(tmp_path, "a.sqlite")
    data = workbook_bytes(tmp_path, a, sample(a))
    b = new_repo(tmp_path, "b.sqlite")

    def boom(*args, **kwargs):
        raise ValueError("disk said no")
    monkeypatch.setattr(Repo, "replace_rows", boom)
    with pytest.raises(ImportProblem) as e:
        import_workbook(b, data, "x.xlsx", "k")
    assert e.value.problems == ["disk said no"]
    assert b.list_workpacks() == []
    assert not any(a == "workpack_imported" for a, _ in b.events)


def test_many_problems_are_capped_on_screen():
    found = ImportProblem([f"p{i}" for i in range(30)]).problems
    assert len(found) == importer.MAX_SHOWN + 1 and found[-1] == "... and 10 more"


# ---- the page ------------------------------------------------------------------------------------
@pytest.fixture()
def client(tmp_path):
    app = create_app(userdata=tmp_path, user_name="tester", modules=["workpacks"])
    return app.test_client(), tmp_path, app


def upload(c, data, name="in.xlsx"):
    return c.post("/workpacks/import", data={"workbook": (io.BytesIO(data), name)}, content_type="multipart/form-data")


def audit(tmp):
    return [a for a, _ in sqlite3.connect(tmp / "audit.sqlite").execute("SELECT action, detail FROM audit")]


def test_importing_through_the_page(client):
    c, tmp, _ = client
    src = new_repo(tmp, "src.sqlite")
    data = workbook_bytes(tmp, src, sample(src))
    r = upload(c, data)
    assert r.status_code == 302 and r.headers["Location"].endswith("/workpacks/1?imported=1")
    html = c.get(r.headers["Location"]).get_data(as_text=True)
    assert "Imported from Excel" in html and "<h1>Wellbay 18 clearance</h1>" in html and "SL1-1" in html
    assert "Imported from Excel" not in c.get("/workpacks/1").get_data(as_text=True)
    assert "workpack_imported" in audit(tmp)


def test_a_refused_workbook_lists_the_problems_and_saves_nothing(client):
    c, tmp, _ = client

    def mutate(wb):
        certs_sheet(wb, ("A", "AS1"))
        wb["Tags"].cell(3, 1).value = "<b>X</b>"
    r = upload(c, make_book(tmp, mutate))
    html = r.get_data(as_text=True)
    assert r.status_code == 422 and "Nothing has been saved" in html
    assert "&lt;b&gt;X&lt;/b&gt;" in html and "<b>X</b>" not in html
    assert "No workpacks yet" in html
    assert "workpack_imported" not in audit(tmp)


def test_no_file_and_wrong_files(client):
    c, tmp, _ = client
    assert c.post("/workpacks/import", data={}, content_type="multipart/form-data").status_code == 400
    assert upload(c, b"not excel").status_code == 422
    assert upload(c, b"x", "notes.txt").status_code == 422


def test_the_import_form_is_only_for_editors(client):
    c, tmp, app = client
    data = make_book(tmp, lambda wb: certs_sheet(wb, ("A", "AS1")))
    assert "Import from Excel" in c.get("/workpacks/").get_data(as_text=True)
    app.config["CTX"]._permissions = {"workpacks.view"}
    assert "Import from Excel" not in c.get("/workpacks/").get_data(as_text=True)
    assert upload(c, data).status_code == 403