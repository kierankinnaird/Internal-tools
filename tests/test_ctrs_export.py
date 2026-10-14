"""The CTR as an Excel file: the layout of the company's CTR sheets, live formulas with their results stored, the logo."""
import datetime
import re
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

import openpyxl
import pytest

from core.migrations import migrate
from modules.ctrs import calc, export
from modules.ctrs.repo import Repo

MIGRATIONS = Path(__file__).resolve().parent.parent / "modules" / "ctrs" / "migrations"


@pytest.fixture
def repo(tmp_path):
    migrate(tmp_path / "ctrs.sqlite", "ctrs", MIGRATIONS)
    return Repo(tmp_path / "ctrs.sqlite")


def payload(**kw):
    body = {
        "title": "Construction - Piping", "ctr_no": "CTR-014", "client": "Apache", "project_name": "Support for Beryl Bravo",
        "project_title": "BB Diesel mods", "ctr_date": "2026-09-17", "framework_ref": "FA-9", "prepared_by": "M.Watt",
        "checked_by": "S.Smith", "approved_by": "R.Craig", "objectives": "Objectives text", "inputs": "Inputs text",
        "scope_notes": "Scope notes", "assumptions": "Assumes", "exclusions": "Excludes", "commencement": "ref plan",
        "completion": "ref plan 2", "duration": "6 weeks", "expenses": "Reimbursable", "payment_terms": "Budget",
        "basis": "1. VAT applies.",
        "lines": [
            {"kind": "allowance", "description": "Progress, Plan Review", "flag": "A", "qty": "15"},
            {"description": "P/Fitter", "grp": "P/Fitter", "eng_override": "137", "from_snapshot": 1},
            {"description": "Fabrication Isometric", "flag": "D", "wbs": "1.2", "qty": "17", "eng_norm": "3", "des_norm": "15", "factor": "1"},
            {"description": "Supports", "flag": "D", "qty": "20", "eng_norm": "5", "des_norm": "15", "factor": "0.5"}],
        "resources": [
            {"name": "Lead Pipefitter", "source": "group:P/Fitter", "split_pct": "60", "rate": "55.5"},
            {"name": "Senior Pipefitter", "source": "group:P/Fitter", "split_pct": "40", "rate": "48"},
            {"name": "Designer", "source": "pool2", "split_pct": "100", "rate": "60"},
            {"name": "Manager", "source": "fixed", "fixed_hours": "12", "rate": "100"}],
        "additions": [{"name": "Offshore Allowance", "unit_cost": "100", "units": "3"}]}
    body.update(kw)
    return body


@pytest.fixture
def built(repo, tmp_path):
    cid = repo.create_ctr({"title": "x"}, "k")
    repo.save_ctr(cid, payload(), "k")
    ctr = repo.get_ctr(cid)
    ctr["source_text"] = "WP1 - Issued for CTR rev A (2026-10-14)"
    path = tmp_path / "out.xlsx"
    export.write_ctr(path, ctr)
    return ctr, path, calc.compute(ctr)


def row_of(ws, label, col=1, start=1):
    for r in range(start, ws.max_row + 1):
        if ws.cell(r, col).value == label:
            return r
    raise AssertionError(f"{label!r} not found in column {col}")


def cached(path):
    """Every stored formula result in the sheet: {'I21': 103.0}."""
    xml = zipfile.ZipFile(path).read("xl/worksheets/sheet1.xml").decode("utf-8")
    return {m.group(1): float(m.group(2)) for m in re.finditer(r'<c r="([A-Z]+\d+)"[^>]*><f>[^<]*</f><v>([^<]*)</v>', xml)}


# ---- the header block ---------------------------------------------------------------------------------------------
def test_the_header_block_is_laid_out_like_the_sheets(built):
    ctr, path, _ = built
    ws = openpyxl.load_workbook(path).active
    assert ws.title == "Construction - Piping"
    assert [ws.cell(r, 1).value for r in (5, 6, 7, 8, 9)] == ["Project Name", "Project Title", "Client", "Date", "Framework Agreement Ref"]
    assert [ws["B5"].value, ws["B6"].value, ws["B7"].value, ws["C9"].value] == ["Support for Beryl Bravo", "BB Diesel mods", "Apache", "FA-9"]
    assert ws["A11"].value == "CTR Title" and ws["B11"].value == "Construction - Piping"
    assert ws["H11"].value == "CTR No" and ws["I11"].value == "CTR-014"
    assert [ws["A13"].value, ws["C13"].value, ws["H13"].value] == ["Prepared by", "Checked by", "Approved by"]
    assert [ws["A14"].value, ws["C14"].value, ws["H14"].value] == ["M.Watt", "S.Smith", "R.Craig"]
    assert ws["A15"].value == "Objectives" and ws["A16"].value == "Objectives text"
    assert ws["A17"].value == "Input Requirements" and ws["A18"].value == "Inputs text"
    assert ws["A19"].value == "Scope of Work"


def test_the_date_is_a_real_date_shown_day_month_year(built):
    ws = openpyxl.load_workbook(built[1]).active
    assert ws["B8"].value == datetime.datetime(2026, 9, 17) and ws["B8"].number_format == "dd/mm/yyyy"


def test_a_date_that_is_not_a_date_stays_as_typed(repo, tmp_path):
    cid = repo.create_ctr({"title": "x"}, "k")
    repo.save_ctr(cid, payload(ctr_date="TBC"), "k")
    export.write_ctr(tmp_path / "t.xlsx", repo.get_ctr(cid))
    assert openpyxl.load_workbook(tmp_path / "t.xlsx").active["B8"].value == "TBC"


def test_the_labels_are_white_on_black_and_the_logo_block_is_merged(built):
    ws = openpyxl.load_workbook(built[1]).active
    assert ws["A5"].fill.fgColor.rgb.endswith("000000") and ws["A5"].font.color.rgb.endswith("FFFFFF") and ws["A5"].font.b
    assert {"A1:D4", "E1:J4", "B5:J5", "B11:E11", "I11:J11", "A15:J15"} <= {str(m) for m in ws.merged_cells.ranges}
    assert ws["A5"].font.name == "Aptos"


# ---- the scope table ------------------------------------------------------------------------------------------
def test_the_scope_table_has_the_sheets_columns_and_the_ctrs_own_column_names(built, repo, tmp_path):
    ctr, path, _ = built
    ws = openpyxl.load_workbook(path).active
    head = row_of(ws, "This CTR covers the following key deliverables / activities:")
    assert [ws.cell(head, c).value for c in (6, 7, 8, 9, 10)] == ["Activity/ Deliverable", "WBS", "Qty", "Eng", "Design"]
    assert [ws.cell(head, c).value for c in (11, 12, 13)] == ["Factor", "Eng NORM", "Design NORM"]
    cid = repo.create_ctr({"title": "y"}, "k")
    repo.save_ctr(cid, payload(pool1_label="Labour", pool2_label="Supervision"), "k")
    export.write_ctr(tmp_path / "r.xlsx", repo.get_ctr(cid))
    ws = openpyxl.load_workbook(tmp_path / "r.xlsx").active
    assert [ws.cell(head, c).value for c in (9, 10, 12, 13)] == ["Labour", "Supervision", "Labour NORM", "Supervision NORM"]


def test_each_line_shows_its_description_flag_wbs_qty_norms_and_factor(built):
    ws = openpyxl.load_workbook(built[1]).active
    r = row_of(ws, "Fabrication Isometric")
    assert [ws.cell(r, c).value for c in (6, 7, 8, 11, 12, 13)] == ["D", "1.2", 17, 1, 3, 15]


def test_line_hours_are_formulas_that_round_up_and_have_their_results_stored(built):
    ctr, path, calced = built
    ws = openpyxl.load_workbook(path).active
    r = row_of(ws, "Fabrication Isometric")
    assert ws.cell(r, 9).value == f'=IF(OR(H{r}="",L{r}=""),0,ROUNDUP(ROUND(H{r}*L{r}*IF(K{r}="",1,K{r}),6),0))'
    assert ws.cell(r, 10).value == f'=IF(OR(H{r}="",M{r}=""),0,ROUNDUP(ROUND(H{r}*M{r}*IF(K{r}="",1,K{r}),6),0))'
    stored = cached(path)
    assert (stored[f"I{r}"], stored[f"J{r}"]) == (51, 255)
    r2 = row_of(ws, "Supports")
    assert (stored[f"I{r2}"], stored[f"J{r2}"]) == (50, 150)                          # 20 x 5 x 0.5 and 20 x 15 x 0.5


def test_typed_hours_are_plain_numbers_not_formulas(built):
    ws = openpyxl.load_workbook(built[1]).active
    r = row_of(ws, "P/Fitter")
    assert ws.cell(r, 9).value == 137 and ws.cell(r, 9).number_format == "0"
    # the other column has no hours typed, so it stays a formula (nothing to calculate yet, so 0)
    assert ws.cell(r, 10).value.startswith("=IF(OR(H") and cached(built[1])[f"J{r}"] == 0


def test_an_allowance_is_a_percentage_cell_and_a_formula_over_the_other_lines_only(built):
    ctr, path, calced = built
    ws = openpyxl.load_workbook(path).active
    r = row_of(ws, "Progress, Plan Review")
    assert ws.cell(r, 8).value == 0.15 and ws.cell(r, 8).number_format == "0%"
    assert ws.cell(r, 9).value == f"=ROUNDUP(ROUND(SUM(I{r + 1}:J{r + 3})*H{r},6),0)"       # not including its own row
    assert cached(path)[f"I{r}"] == calced["lines"][0]["eng"] == 97                           # 15% of (137 + 306 + 200) = 96.45
    assert ws.cell(r, 10).value == 0


def test_an_allowance_in_the_middle_skips_its_own_row_in_the_sum(repo, tmp_path):
    cid = repo.create_ctr({"title": "x"}, "k")
    repo.save_ctr(cid, payload(lines=[{"description": "a", "eng_override": "100"}, {"kind": "allowance", "qty": "10"},
                                      {"description": "b", "eng_override": "50"}, {"kind": "allowance", "qty": "20"}]), "k")
    export.write_ctr(tmp_path / "t.xlsx", repo.get_ctr(cid))
    ws = openpyxl.load_workbook(tmp_path / "t.xlsx").active
    first = row_of(ws, "a")
    assert ws.cell(first + 1, 9).value == f"=ROUNDUP(ROUND(SUM(I{first}:J{first},I{first + 2}:J{first + 2})*H{first + 1},6),0)"
    assert cached(tmp_path / "t.xlsx")[f"I{first + 1}"] == 15 and cached(tmp_path / "t.xlsx")[f"I{first + 3}"] == 30


def test_the_jobcards_source_is_stated_under_the_table_and_the_page_breaks_after_it(built):
    ctr, path, _ = built
    ws = openpyxl.load_workbook(path).active
    r = row_of(ws, "Hours from the jobcards: WP1 - Issued for CTR rev A (2026-10-14)")
    assert ws.cell(r, 1).font.i
    assert [b.id for b in ws.row_breaks.brk] == [r]
    assert ws.cell(r + 1, 1).value == "Scope notes"


def test_without_a_jobcards_source_there_is_no_source_row(repo, tmp_path):
    cid = repo.create_ctr({"title": "x"}, "k")
    repo.save_ctr(cid, payload(), "k")
    export.write_ctr(tmp_path / "t.xlsx", repo.get_ctr(cid))
    ws = openpyxl.load_workbook(tmp_path / "t.xlsx").active
    assert all("Hours from the jobcards" not in str(ws.cell(r, 1).value) for r in range(1, ws.max_row + 1))


def test_the_three_text_blocks_appear_only_when_there_is_something_to_say(repo, tmp_path):
    cid = repo.create_ctr({"title": "x"}, "k")
    repo.save_ctr(cid, payload(scope_notes="", assumptions="Only this", exclusions="  "), "k")
    export.write_ctr(tmp_path / "t.xlsx", repo.get_ctr(cid))
    ws = openpyxl.load_workbook(tmp_path / "t.xlsx").active
    values = [ws.cell(r, 1).value for r in range(1, ws.max_row + 1)]
    assert "Only this" in values and "Scope notes" not in values and "Excludes" not in values


# ---- resources, additions, totals ---------------------------------------------------------------------------------
def test_each_resource_has_its_hours_rate_and_cost_formula(built):
    ctr, path, calced = built
    ws = openpyxl.load_workbook(path).active
    head = row_of(ws, "Resource")
    assert [ws.cell(head, c).value for c in (1, 3, 5, 9)] == ["Resource", "Cost £ / Hour", "Hours", "Total"]
    lead = row_of(ws, "Lead Pipefitter")
    assert [ws.cell(lead, 3).value, ws.cell(lead, 5).value, ws.cell(lead, 9).value] == [55.5, 82, f"=E{lead}*C{lead}"]
    senior = row_of(ws, "Senior Pipefitter")
    assert [ws.cell(senior, 3).value, ws.cell(senior, 5).value] == [48, 55]                    # 82 + 55 = 137, as the jobcards
    stored = cached(path)
    assert stored[f"I{lead}"] == round(82 * 55.5, 2) and stored[f"I{senior}"] == 55 * 48
    assert ws.cell(lead, 3).number_format == "[$£-809]#,##0.00"


def test_the_split_and_where_the_hours_come_from_sit_beside_the_table_outside_the_print_area(built):
    ws = openpyxl.load_workbook(built[1]).active
    head = row_of(ws, "Resource")
    assert [ws.cell(head, 11).value, ws.cell(head, 12).value] == ["Split", "Takes hours from"]
    lead, senior, designer, manager = (row_of(ws, n) for n in ("Lead Pipefitter", "Senior Pipefitter", "Designer", "Manager"))
    assert (ws.cell(lead, 11).value, ws.cell(senior, 11).value, ws.cell(designer, 11).value) == (0.6, 0.4, 1.0)
    assert ws.cell(lead, 11).number_format == "0.##%"
    assert ws.cell(lead, 12).value == "“P/Fitter” hours" and ws.cell(designer, 12).value == "Design hours"
    assert ws.cell(manager, 11).value is None and ws.cell(manager, 12).value == "typed hours"
    assert ws.cell(manager, 5).value == 12


def test_the_totals_are_formulas_with_their_results_stored(built):
    ctr, path, calced = built
    ws = openpyxl.load_workbook(path).active
    first, last = row_of(ws, "Lead Pipefitter"), row_of(ws, "Resource") + 4
    total = last + 1
    assert ws.cell(total, 5).value == f"=SUM(E{first}:E{last})" and ws.cell(total, 9).value == f"=SUM(I{first}:I{last})"
    stored = cached(path)
    assert stored[f"E{total}"] == calced["totals"]["hours"] and stored[f"I{total}"] == calced["totals"]["resources_cost"]
    assert ws.cell(total, 5).fill.fgColor.rgb.endswith("92D050")


def test_the_additions_and_the_grand_total(built):
    ctr, path, calced = built
    ws = openpyxl.load_workbook(path).active
    head = row_of(ws, "Addition")
    assert [ws.cell(head, c).value for c in (3, 5, 9)] == ["Cost £ / Unit", "Unit", "Total"]
    r = row_of(ws, "Offshore Allowance")
    assert [ws.cell(r, 3).value, ws.cell(r, 5).value, ws.cell(r, 9).value] == [100, 3, f"=C{r}*E{r}"]
    grand = head + 3                                    # two addition rows (the table always shows at least two), then the total
    assert ws.cell(grand, 9).value == f"=SUM(I{r}:I{r + 1})+I{head - 1}"
    assert cached(path)[f"I{grand}"] == calced["totals"]["cost"] and ws.cell(grand, 9).fill.fgColor.rgb.endswith("92D050")


def test_the_tables_always_show_at_least_four_resource_and_two_addition_rows(repo, tmp_path):
    cid = repo.create_ctr({"title": "x"}, "k")
    export.write_ctr(tmp_path / "t.xlsx", repo.get_ctr(cid))
    ws = openpyxl.load_workbook(tmp_path / "t.xlsx").active
    head = row_of(ws, "Resource")
    assert ws.cell(head + 5, 5).value == "=SUM(E{0}:E{1})".format(head + 1, head + 4)
    assert ws.cell(head + 6, 1).value == "Addition"
    assert ws.cell(head + 7, 1).value == "Offshore Allowance" and ws.cell(head + 8, 1).value == "Mob / De-Mob Travel Allowance"


@pytest.mark.parametrize("terms, cell", [("Reimbursable", "D"), ("Budget", "H"), ("Lump Sum", "J")])
def test_the_payment_terms_tick_sits_in_the_right_box(repo, tmp_path, terms, cell):
    cid = repo.create_ctr({"title": "x"}, "k")
    repo.save_ctr(cid, payload(payment_terms=terms), "k")
    export.write_ctr(tmp_path / "t.xlsx", repo.get_ctr(cid))
    ws = openpyxl.load_workbook(tmp_path / "t.xlsx").active
    r = row_of(ws, "Payment Terms")
    assert [c for c in "DHJ" if ws[f"{c}{r}"].value == "X"] == [cell]
    assert [ws[f"{c}{r}"].value for c in ("B", "E", "I")] == ["Reimbursable", "Budget", "Lump Sum"]


def test_the_dates_duration_expenses_and_basis_are_there(built):
    ws = openpyxl.load_workbook(built[1]).active
    r = row_of(ws, "Commencement Date")
    assert [ws.cell(r, 2).value, ws.cell(r, 8).value, ws.cell(r, 9).value] == ["ref plan", "Duration", "6 weeks"]
    assert [ws.cell(r + 1, 1).value, ws.cell(r + 1, 2).value, ws.cell(r + 1, 9).value] == ["Completion Date", "ref plan 2", "Reimbursable"]
    b = row_of(ws, "Basis of Proposal")
    assert ws.cell(b + 1, 1).value == "1. VAT applies."


# ---- the file itself ------------------------------------------------------------------------------------------
def test_every_formula_has_its_result_stored_beside_it(built):
    """What Excel's Protected View shows: no calculation happens there, so a formula without a stored result is blank."""
    ctr, path, _ = built
    ws = openpyxl.load_workbook(path).active
    formulas = {c.coordinate for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")}
    assert formulas and formulas <= set(cached(path))


def test_the_stored_results_agree_with_the_tools_own_figures(built):
    ctr, path, calced = built
    ws = openpyxl.load_workbook(path).active
    stored = cached(path)
    lines_first = row_of(ws, "Progress, Plan Review")
    for i, res in enumerate(calced["lines"]):
        for col, key in (("I", "eng"), ("J", "des")):
            ref = f"{col}{lines_first + i}"
            if ref in stored:
                assert stored[ref] == res[key], ref
    resource_first = row_of(ws, "Lead Pipefitter")
    for i, res in enumerate(calced["resources"]):
        assert stored[f"I{resource_first + i}"] == res["cost"]


def test_the_logo_is_in_the_file_and_the_file_is_well_formed(built):
    ctr, path, _ = built
    z = zipfile.ZipFile(path)
    assert z.testzip() is None
    assert z.read("xl/media/image1.png") == export.LOGO.read_bytes() and z.read("xl/media/image1.png")[:8] == b"\x89PNG\r\n\x1a\n"
    names = z.namelist()
    assert {"xl/drawings/drawing1.xml", "xl/drawings/_rels/drawing1.xml.rels", "xl/worksheets/_rels/sheet1.xml.rels"} <= set(names)
    assert b'Extension="png"' in z.read("[Content_Types].xml") and b"/xl/drawings/drawing1.xml" in z.read("[Content_Types].xml")
    sheet = z.read("xl/worksheets/sheet1.xml").decode()
    assert sheet.count("<drawing ") == 1 and sheet.rstrip().endswith("</worksheet>")
    for name in names:
        if name.endswith((".xml", ".rels")):
            ET.fromstring(z.read(name))                                   # raises if any part is not well-formed
    assert len(names) == len(set(names))                                   # no part written twice


def test_the_file_opens_again_in_openpyxl_with_the_logo_part_in_place(built):
    wb = openpyxl.load_workbook(built[1])
    assert wb.sheetnames == ["Construction - Piping"]


def test_a_missing_logo_file_still_gives_a_good_workbook(built, monkeypatch, tmp_path):
    ctr, _, calced = built
    monkeypatch.setattr(export, "LOGO", tmp_path / "nope.png")
    export.write_ctr(tmp_path / "nologo.xlsx", ctr, calced)
    z = zipfile.ZipFile(tmp_path / "nologo.xlsx")
    assert "xl/media/image1.png" not in z.namelist() and b"<drawing" not in z.read("xl/worksheets/sheet1.xml")
    assert cached(tmp_path / "nologo.xlsx")                                # results are still stored
    assert openpyxl.load_workbook(tmp_path / "nologo.xlsx").active["B5"].value == "Support for Beryl Bravo"


def test_it_is_one_sheet_set_up_for_an_a4_portrait_page_at_full_width(built):
    ctr, path, _ = built
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    assert len(wb.sheetnames) == 1
    assert int(ws.page_setup.paperSize) == int(ws.PAPERSIZE_A4) == 9 and ws.page_setup.orientation == "portrait"
    assert ws.page_setup.fitToWidth == 1 and ws.page_setup.fitToHeight == 0 and ws.sheet_properties.pageSetUpPr.fitToPage
    last = row_of(ws, "Basis of Proposal") + 1
    assert ws.print_area.endswith(f"$A$1:$J${last}")                      # K to M are working figures, not printed


def test_column_widths_follow_the_sheets(built):
    ws = openpyxl.load_workbook(built[1]).active
    assert round(ws.column_dimensions["A"].width, 1) == 22.3 and round(ws.column_dimensions["G"].width, 1) == 16.3


# ---- awkward text -------------------------------------------------------------------------------------------
def test_typed_text_that_starts_with_an_equals_sign_stays_text(repo, tmp_path):
    cid = repo.create_ctr({"title": "x"}, "k")
    repo.save_ctr(cid, payload(objectives="=SUM(1,2)", client="=cmd|' /C calc'!A0",
                               lines=[{"description": "=HYPERLINK(\"http://evil\")", "eng_override": "1"}]), "k")
    export.write_ctr(tmp_path / "t.xlsx", repo.get_ctr(cid))
    ws = openpyxl.load_workbook(tmp_path / "t.xlsx").active
    for ref in ("A16", "B7", "A21"):
        assert ws[ref].data_type == "s", ref
    assert ws["A16"].value == "=SUM(1,2)"
    assert "<f>SUM(1,2)</f>" not in zipfile.ZipFile(tmp_path / "t.xlsx").read("xl/worksheets/sheet1.xml").decode()


def test_control_characters_are_removed_so_excel_does_not_call_the_file_corrupt():
    assert export.text("a\x00b\x07c\x1fd") == "abcd" and export.text(None) == "" and export.text("ok\nline") == "ok\nline"


def test_long_text_gets_a_taller_row(repo, tmp_path):
    cid = repo.create_ctr({"title": "x"}, "k")
    repo.save_ctr(cid, payload(objectives="word " * 400, assumptions="short"), "k")
    export.write_ctr(tmp_path / "t.xlsx", repo.get_ctr(cid))
    ws = openpyxl.load_workbook(tmp_path / "t.xlsx").active
    assert ws.row_dimensions[16].height > 200 and ws.row_dimensions[18].height == 45


def test_a_long_scope_description_gets_a_taller_row(repo, tmp_path):
    cid = repo.create_ctr({"title": "x"}, "k")
    repo.save_ctr(cid, payload(lines=[{"description": "A long activity " * 12, "eng_override": "1"}, {"description": "Short", "eng_override": "1"}]), "k")
    export.write_ctr(tmp_path / "t.xlsx", repo.get_ctr(cid))
    ws = openpyxl.load_workbook(tmp_path / "t.xlsx").active
    long_row, short_row = row_of(ws, ("A long activity " * 12).strip()), row_of(ws, "Short")
    assert ws.row_dimensions[long_row].height > ws.row_dimensions[short_row].height


@pytest.mark.parametrize("title, expected", [("Piping", "Piping"), ("A/B:C*D?[E]\\F", "A B C D  E  F"), ("", "CTR"), ("   ", "CTR"),
                                             ("x" * 50, "x" * 31), ("[:]", "CTR")])
def test_the_sheet_name_is_always_one_excel_accepts(title, expected):
    assert export.sheet_name(title) == expected


def test_an_empty_ctr_exports_without_error(repo, tmp_path):
    cid = repo.create_ctr({"title": "Empty"}, "k")
    export.write_ctr(tmp_path / "t.xlsx", repo.get_ctr(cid))
    ws = openpyxl.load_workbook(tmp_path / "t.xlsx").active
    assert ws["B11"].value == "Empty" and zipfile.ZipFile(tmp_path / "t.xlsx").testzip() is None


def test_an_empty_scope_table_still_shows_one_row_to_fill_in(repo, tmp_path):
    cid = repo.create_ctr({"title": "Empty"}, "k")
    export.write_ctr(tmp_path / "t.xlsx", repo.get_ctr(cid))
    ws = openpyxl.load_workbook(tmp_path / "t.xlsx").active
    head = row_of(ws, "This CTR covers the following key deliverables / activities:")
    assert {str(m) for m in ws.merged_cells.ranges} >= {f"A{head + 1}:E{head + 1}"}
    assert [b.id for b in ws.row_breaks.brk] == [head + 1]


def test_a_big_ctr_exports(repo, tmp_path):
    cid = repo.create_ctr({"title": "Big"}, "k")
    repo.save_ctr(cid, payload(lines=[{"description": f"line {i}", "qty": "1", "eng_norm": "2"} for i in range(400)],
                               resources=[{"name": f"r{i}", "source": "pool1", "split_pct": "1", "rate": "1"} for i in range(60)],
                               additions=[{"name": f"a{i}", "unit_cost": "1", "units": "1"} for i in range(30)]), "k")
    export.write_ctr(tmp_path / "t.xlsx", repo.get_ctr(cid))
    assert openpyxl.load_workbook(tmp_path / "t.xlsx").active["A1"].value is None
    assert cached(tmp_path / "t.xlsx")
