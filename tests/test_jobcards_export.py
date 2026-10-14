from pathlib import Path

import openpyxl
import pytest
from openpyxl.cell.rich_text import CellRichText, TextBlock

from core.migrations import migrate
from modules.jobcards import export
from modules.jobcards.link import JobcardsLink
from modules.jobcards.repo import Repo

MIGRATIONS = Path(__file__).resolve().parent.parent / "modules" / "jobcards" / "migrations"
FIRST = "Raise PTW, Carry out TBT, Erect Signs & Barriers"
LAST = "Complete this jobcard"


@pytest.fixture()
def repo(tmp_path):
    db = tmp_path / "jobcards.sqlite"
    migrate(db, "jobcards", MIGRATIONS)
    return Repo(db)


def make(repo, tasks=(), **fields):
    wp = repo.create_workpack({"number": "BB-26-M0001", "title": "BB-18 WELLBAY CLEARANCE - PHASE 2", "fcp_no": "FCP-1",
                               "jobcard_prefix": "NMC-BB18-"}, "k")
    jc_id = repo.create_jobcard(wp, "k")
    jc = repo.get_jobcard(jc_id)
    first, last = jc["tasks"]
    body = dict({"title": "Clear wellbay", "trade": "P/Fitter", "revision": "B", "asset": "Wellbay 18", "ans_sh1_no": "SH1-9",
                 "created_by_name": "A Person", "created_by_position": "Engineer", "created_by_signature": "AP",
                 "created_date": "2026-10-14", "additional_info": "Line one\nLine two",
                 "flags": ["COSHH", "Rope access", "Operations"],
                 "drawings": [{"number": "D-1", "title": "General arrangement"}],
                 "tasks": [dict(first, id=first["id"]), *tasks, dict(last, id=last["id"])]}, **fields)
    saved = repo.save_jobcard(jc_id, body, "k")
    return saved, repo.get_workpack(wp)


def build(tmp_path, jc, wp, rich=True):
    path = tmp_path / "out.xlsx"
    export.write_jobcard(path, jc, wp)
    return openpyxl.load_workbook(path, rich_text=rich)


def find(ws, label, col=1):
    for row in range(1, ws.max_row + 1):
        if ws.cell(row, col).value == label:
            return row
    raise AssertionError(f"{label!r} not found")


def mine(description, **more):
    return dict({"description": description, "qty": 2, "norm_hours": 3, "source": "manual", "basis": "quote"}, **more)


# ---- page one ---------------------------------------------------------------------------------
def test_it_is_one_sheet_with_a_page_break_before_the_tasks(repo, tmp_path):
    """Excel prints only the active sheet, so the tasks must be on the same sheet as page 1, after a page break."""
    wb = build(tmp_path, *make(repo))
    assert wb.sheetnames == ["Jobcard"]
    ws = wb["Jobcard"]
    (page_break,) = [b.id for b in ws.row_breaks.brk]
    assert find(ws, "CREATED BY") < page_break < find(ws, "TASKS") < find(ws, "ITEM")
    assert ws.cell(page_break + 1, 4).value == "CONSTRUCTION JOBCARD"             # page 2 starts with the header again


def test_the_header_block_is_filled_at_the_top_of_both_pages(repo, tmp_path):
    jc, wp = make(repo)
    ws = build(tmp_path, jc, wp)["Jobcard"]
    tops = [row for row in range(1, ws.max_row + 1) if ws.cell(row, 4).value == "CONSTRUCTION JOBCARD"]
    assert len(tops) == 2 and tops[0] == 1
    for top in tops:
        assert ws.cell(top, 1).value == "NMC ENERGY"
        assert (ws.cell(top + 1, 3).value, ws.cell(top + 1, 7).value, ws.cell(top + 1, 10).value) == ("NMC-BB18-001", "B", "FCP-1")
        assert ws.cell(top + 2, 3).value == "Clear wellbay"
        assert (ws.cell(top + 4, 1).value, ws.cell(top + 4, 4).value) == ("BB-26-M0001", "BB-18 WELLBAY CLEARANCE - PHASE 2")
        assert (ws.cell(top + 6, 1).value, ws.cell(top + 6, 4).value, ws.cell(top + 6, 7).value) == ("Wellbay 18", "SH1-9", "P/Fitter")


def test_page_one_has_the_people_information_flags_and_drawings(repo, tmp_path):
    ws = build(tmp_path, *make(repo))["Jobcard"]
    r = find(ws, "CREATED BY")
    assert (ws.cell(r + 1, 2).value, ws.cell(r + 1, 8).value) == ("A Person", "Engineer")
    assert (ws.cell(r + 2, 2).value, ws.cell(r + 2, 8).value) == ("AP", "2026-10-14")
    r = find(ws, "Additional information")
    assert ws.cell(r + 1, 1).value == "Line one\nLine two"
    ticked = {ws.cell(row, col).value: ws.cell(row, tick).value
              for row in range(r + 2, r + 9) for col, tick in ((1, 3), (4, 6), (7, 11))}
    assert ticked["COSHH"] == "✓" and ticked["Rope access"] == "✓" and ticked["Operations"] == "✓"
    assert ticked["Cold work"] is None and ticked["Subsea"] is None and len(ticked) == 21
    r = find(ws, "DRAWING REFERENCE")
    assert (ws.cell(r + 2, 1).value, ws.cell(r + 2, 4).value) == ("D-1", "General arrangement")
    assert ws.cell(r + 3, 1).value in (None, "") and ws.cell(r + 4, 1).value in (None, "")        # at least 3 rows
    assert ws.cell(find(ws, "JOBCARD COMPLETION") + 1, 1).value == "COMPLETED BY"


def test_a_blank_card_still_builds(repo, tmp_path):
    wp = repo.create_workpack({"number": "W"}, "k")
    jc = repo.get_jobcard(repo.create_jobcard(wp, "k"))
    wb = build(tmp_path, jc, repo.get_workpack(wp))
    assert wb["Jobcard"]["C2"].value == "W-001"


# ---- the tasks --------------------------------------------------------------------------------
def test_tasks_run_first_standard_to_closing_standard_with_formulas(repo, tmp_path):
    ws = build(tmp_path, *make(repo, [mine("Cut it"), mine("Weld it", qty=4, norm_hours=1.5)]))["Jobcard"]
    head = find(ws, "ITEM")
    rows = [head + i for i in range(1, 11)]
    assert [ws.cell(r, 1).value for r in rows] == list(range(1, 11))                    # at least 10 item rows
    assert ws.cell(rows[0], 2).value == FIRST and ws.cell(rows[1], 2).value == "Cut it" and ws.cell(rows[2], 2).value == "Weld it"
    assert ws.cell(rows[-1], 2).value == LAST                                           # the closing task is on the bottom row
    assert all(ws.cell(r, 2).value in (None, "") for r in rows[3:-1])                   # blank rows sit above it
    assert [ws.cell(rows[2], c).value for c in (8, 9, 10)] == [4, 1.5, 1]
    assert [ws.cell(rows[0], c).value for c in (8, 9, 10)] == [1, 1, 1]
    r = rows[1]
    assert ws.cell(r, 11).value == f'=IF(OR(H{r}="",I{r}=""),0,ROUNDUP(ROUND(H{r}*I{r}*IF(J{r}="",1,J{r}),6),0))'
    total = rows[-1] + 1
    assert ws.cell(total, 1).value == "TOTAL MANHOURS:" and ws.cell(total, 11).value == f"=SUM(K{rows[0]}:K{rows[-1]})"


def test_a_long_list_gets_a_row_for_every_task_and_the_closing_one_is_last(repo, tmp_path):
    ws = build(tmp_path, *make(repo, [mine(f"task {i}") for i in range(1, 14)]))["Jobcard"]
    head = find(ws, "ITEM")
    descriptions = [ws.cell(head + i, 2).value for i in range(1, 16)]
    assert descriptions[0] == FIRST and descriptions[1] == "task 1" and descriptions[13] == "task 13" and descriptions[14] == LAST
    assert ws.cell(head + 16, 1).value == "TOTAL MANHOURS:"


def test_a_factor_reason_is_shown_under_the_description(repo, tmp_path):
    ws = build(tmp_path, *make(repo, [mine("Tight spot", factor=1.5, factor_reason="Congested")]))["Jobcard"]
    cell = ws.cell(find(ws, "ITEM") + 2, 2).value
    assert isinstance(cell, CellRichText) and str(cell) == "Tight spot\nFactor: Congested"
    assert ws.cell(find(ws, "ITEM") + 2, 10).value == 1.5


def test_a_note_is_a_heading_with_no_hours_and_the_subtasks_under_it_carry_them(repo, tmp_path):
    heading = "<b>Destruct hydraulic supply tubing &amp; Containment from interface plate to:</b>"
    ws = build(tmp_path, *make(repo, [
        mine("", kind="note", note=heading),
        mine("Annulus Kill MSAS Valve", qty=1, norm_hours=2), mine("Downhole Safety Valve", qty=1, norm_hours=3)]))["Jobcard"]
    head = find(ws, "ITEM")
    note_row, first_sub, second_sub = head + 2, head + 3, head + 4
    cell = ws.cell(note_row, 2).value
    assert isinstance(cell, CellRichText) and str(cell) == "Destruct hydraulic supply tubing & Containment from interface plate to:"
    assert [ws.cell(note_row, c).value for c in (8, 9, 10)] == [None, None, None]          # no quantity, norm or factor
    assert ws.cell(first_sub, 2).value == "Annulus Kill MSAS Valve" and (ws.cell(first_sub, 8).value, ws.cell(first_sub, 9).value) == (1, 2)
    assert ws.cell(second_sub, 2).value == "Downhole Safety Valve" and (ws.cell(second_sub, 8).value, ws.cell(second_sub, 9).value) == (1, 3)
    assert ws.cell(note_row, 11).value.startswith("=IF(")                                    # its hours formula gives 0 with no qty


def test_a_note_keeps_its_bold_italic_underline_and_bullets(repo, tmp_path):
    note = ("Destruct to:<ul><li>Annulus Kill MSAS Valve</li><li><b>Downhole Safety Valve</b></li></ul>"
            "<div><i>Check</i> <u>twice</u></div><ol><li>one</li><li>two</li></ol>")
    ws = build(tmp_path, *make(repo, [mine("", kind="note", note=note)]))["Jobcard"]
    cell = ws.cell(find(ws, "ITEM") + 2, 2).value
    assert isinstance(cell, CellRichText)
    assert str(cell) == "Destruct to:\n• Annulus Kill MSAS Valve\n• Downhole Safety Valve\nCheck twice\n1. one\n2. two"
    runs = {block.text.strip(): block.font for block in cell if isinstance(block, TextBlock)}     # a line break rides on the run after it
    assert runs["Downhole Safety Valve"].b and not runs["Annulus Kill MSAS Valve"].b
    assert runs["Check"].i and runs["twice"].u == "single" and not runs["Check"].b
    assert runs["Destruct to:"].rFont == "Calibri" and runs["Destruct to:"].sz == 10


def test_a_note_with_formatting_is_rich_text(repo, tmp_path):
    ws = build(tmp_path, *make(repo, [mine("", kind="note", note="Just <i>text</i> here")]))["Jobcard"]
    assert isinstance(ws.cell(find(ws, "ITEM") + 2, 2).value, CellRichText)


def test_a_note_with_no_formatting_is_a_plain_cell(repo, tmp_path):
    ws = build(tmp_path, *make(repo, [mine("", kind="note", note="Just text")]))["Jobcard"]
    assert ws.cell(find(ws, "ITEM") + 2, 2).value == "Just text"


def test_a_task_is_a_plain_cell(repo, tmp_path):
    ws = build(tmp_path, *make(repo, [mine("Plain")]))["Jobcard"]
    assert ws.cell(find(ws, "ITEM") + 2, 2).value == "Plain"


def test_text_that_starts_with_an_equals_sign_stays_text(repo, tmp_path):
    jc, wp = make(repo, [mine("=1+1")], title="=SUM(A1)")
    wb = build(tmp_path, jc, wp, rich=False)
    assert wb["Jobcard"].cell(find(wb["Jobcard"], "ITEM") + 2, 2).data_type == "s"
    assert wb["Jobcard"].cell(find(wb["Jobcard"], "ITEM") + 2, 2).value == "=1+1"
    assert wb["Jobcard"]["C3"].data_type == "s" and wb["Jobcard"]["C3"].value == "=SUM(A1)"


def test_characters_excel_cannot_hold_are_dropped_not_fatal(repo, tmp_path):
    jc, wp = make(repo, [mine("Bell\x07 and null\x00 char"), mine("", kind="note", note="<b>b\x0bell</b>")], title="T\x01itle")
    ws = build(tmp_path, jc, wp)["Jobcard"]
    assert ws.cell(find(ws, "ITEM") + 2, 2).value == "Bell and null char"
    assert str(ws.cell(find(ws, "ITEM") + 3, 2).value) == "bell"
    assert ws["C3"].value == "Title"


def test_printing_is_set_up_for_a_portrait_a4_page_wide(repo, tmp_path):
    wb = build(tmp_path, *make(repo))
    for ws in wb.worksheets:
        assert int(ws.page_setup.paperSize) == 9 and ws.page_setup.orientation == "portrait"
        assert ws.page_setup.fitToWidth == 1 and ws.page_setup.fitToHeight == 0 and ws.sheet_properties.pageSetUpPr.fitToPage


def test_a_card_without_standard_tasks_exports_as_it_is(repo, tmp_path):
    wp = repo.create_workpack({"number": "W"}, "k")
    jc = repo.get_jobcard(repo.create_jobcard(wp, "k"))
    jc["tasks"] = [{"description": "Only", "qty": 1.0, "norm_hours": 1.0, "factor": 1.0, "role": "", "note": ""}]
    ws = build(tmp_path, jc, repo.get_workpack(wp))["Jobcard"]
    head = find(ws, "ITEM")
    assert ws.cell(head + 1, 2).value == "Only" and all(ws.cell(head + i, 2).value in (None, "") for i in range(2, 11))


# ---- the service: every linked jobcard into a folder ---------------------------------------------
@pytest.fixture()
def link(tmp_path):
    db = tmp_path / "jobcards.sqlite"
    migrate(db, "jobcards", MIGRATIONS)
    return JobcardsLink(db)


def test_export_linked_writes_one_file_per_jobcard_of_every_linked_workpack(link, tmp_path):
    link.create_for(1, "k", {"number": "A", "jobcard_prefix": "NMC-BB18-"}, count=2)
    link.create_for(1, "k", {"number": "B", "jobcard_prefix": "NMC-BB19-"}, count=1)
    link.create_for(2, "k", {"number": "C", "jobcard_prefix": "OTHER-"}, count=3)         # linked elsewhere: not included
    Repo(link.db_path).create_workpack({"number": "D"}, "k")                              # not linked at all
    folder = tmp_path / "pack" / "9.0 JOBCARDS"
    result = link.export_linked(1, folder)
    assert result == [{"number": "NMC-BB18-001", "file": "NMC-BB18-001.xlsx", "error": ""},
                      {"number": "NMC-BB18-002", "file": "NMC-BB18-002.xlsx", "error": ""},
                      {"number": "NMC-BB19-001", "file": "NMC-BB19-001.xlsx", "error": ""}]
    assert sorted(p.name for p in folder.iterdir()) == ["NMC-BB18-001.xlsx", "NMC-BB18-002.xlsx", "NMC-BB19-001.xlsx"]
    assert openpyxl.load_workbook(folder / "NMC-BB19-001.xlsx")["Jobcard"]["C2"].value == "NMC-BB19-001"


def test_nothing_linked_writes_nothing_and_makes_no_folder(link, tmp_path):
    assert link.export_linked(1, tmp_path / "pack" / "9.0 JOBCARDS") == []
    assert not (tmp_path / "pack").exists()
    with pytest.raises(ValueError, match="A workpack is required"):
        link.export_linked(0, tmp_path / "x")


def test_a_number_with_unsafe_characters_becomes_a_safe_file_name(link, tmp_path):
    link.create_for(1, "k", {"number": "A", "jobcard_prefix": 'BB/18:"x"*?-'}, count=1)
    link.create_for(1, "k", {"number": "B", "jobcard_prefix": "BB_18__x___-"}, count=1)      # sanitises to the same name
    result = link.export_linked(1, tmp_path / "out")
    names = [r["file"] for r in result]
    assert names == ["BB_18__x___-001.xlsx", "BB_18__x___-001 (2).xlsx"] and all(r["error"] == "" for r in result)
    assert sorted(p.name for p in (tmp_path / "out").iterdir()) == sorted(names)


def test_one_jobcard_failing_does_not_stop_the_others(link, tmp_path, monkeypatch):
    link.create_for(1, "k", {"number": "A", "jobcard_prefix": "J-"}, count=3)
    real = export.write_jobcard

    def flaky(path, jc, wp):
        if jc["number"] == "J-002":
            path.write_bytes(b"half")
            raise RuntimeError("disk full")
        return real(path, jc, wp)
    monkeypatch.setattr(export, "write_jobcard", flaky)
    result = link.export_linked(1, tmp_path / "out")
    assert [(r["number"], bool(r["file"]), r["error"]) for r in result] == [
        ("J-001", True, ""), ("J-002", False, "disk full"), ("J-003", True, "")]
    assert sorted(p.name for p in (tmp_path / "out").iterdir()) == ["J-001.xlsx", "J-003.xlsx"]     # no half-written file


# ---- what Excel itself needs (found by opening the files in Excel and printing them to PDF) ----------
def sheet_xml(path, sheet):
    import zipfile
    return zipfile.ZipFile(path).read(f"xl/worksheets/sheet{sheet}.xml").decode("utf-8")


def test_no_rich_text_run_is_only_white_space(repo, tmp_path):
    """Excel calls a file corrupt, and repairs it by dropping content, if a run is just a line break."""
    import re
    note = "Head:<ul><li>One</li><li><b>Two</b></li></ul><div><i>Check</i> <u>twice</u></div>"
    jc, wp = make(repo, [mine("Task", factor=1.5, factor_reason="Congested"), mine("", kind="note", note=note),
                         mine("", kind="note", note="<b>x</b>")])
    path = tmp_path / "out.xlsx"
    export.write_jobcard(path, jc, wp)
    xml = sheet_xml(path, 1)
    runs = re.findall(r"<r>.*?<t[^>]*>(.*?)</t></r>", xml, flags=re.S)
    assert runs and all(r.strip() for r in runs)


def test_no_blank_runs_folds_white_space_into_the_next_run():
    plain = (False, False, False)
    assert export._no_blank_runs([("A", *plain), ("\n", *plain), ("B", True, False, False)]) == [("A", *plain), ("\nB", True, False, False)]
    assert export._no_blank_runs([("A", *plain), (" ", *plain), ("B", False, True, False)]) == [("A", *plain), (" B", False, True, False)]
    assert export._no_blank_runs([("A", *plain), ("\n", *plain), (" ", *plain)]) == [("A", *plain)]        # nothing follows: dropped
    assert export._no_blank_runs([("", *plain), ("\n", *plain), ("Note", *plain)]) == [("Note", *plain)]    # no leading blank line
    assert export._no_blank_runs([]) == []


def test_hours_and_total_are_real_formulas_but_typed_text_is_not(repo, tmp_path):
    jc, wp = make(repo, [mine("=1+1")])
    ws = build(tmp_path, jc, wp, rich=False)["Jobcard"]
    head = find(ws, "ITEM")
    assert ws.cell(head + 1, 11).data_type == "f" and ws.cell(head + 10, 11).data_type == "f"    # hours, including blank rows
    assert ws.cell(head + 11, 11).data_type == "f"                                              # the total
    assert ws.cell(head + 2, 2).data_type == "s"                                                # typed "=1+1" stays text


def test_the_sheet_is_as_wide_as_an_a4_page_so_excel_prints_it_full_size():
    pixels = sum(int(w * 7 + 5) for w in export.WIDTHS)                    # Excel's column width to pixels
    assert export.A4_PIXELS * 0.95 <= pixels <= export.A4_PIXELS           # fills the page, and is never shrunk to fit
    assert len(export.WIDTHS) == 11


# ---- the results are stored beside the formulas -----------------------------------------------------
def cached(path):
    """What a viewer that does not calculate sees (Excel's Protected View, a file previewer, a PDF converter)."""
    ws = openpyxl.load_workbook(path, data_only=True)["Jobcard"]
    head = find(ws, "ITEM")
    return ws, head


def test_hours_and_total_are_stored_in_the_file_not_just_as_formulas(repo, tmp_path):
    jc, wp = make(repo, [mine("", kind="note", note="<b>Heading</b>"), mine("Sub one", qty=1, norm_hours=2.5),
                         mine("Sub two", qty=2, norm_hours=3, factor=1.5, factor_reason="Congested")])
    path = tmp_path / "out.xlsx"
    export.write_jobcard(path, jc, wp)
    ws, head = cached(path)
    hours = [ws.cell(head + i, 11).value for i in range(1, 11)]
    #          standard  heading  sub one  sub two  blank x5     closing standard
    assert hours == [1, 0, 3, 9, 0, 0, 0, 0, 0, 1]                            # 2.5 is rounded up to 3 on its own
    assert ws.cell(head + 11, 11).value == 14                                      # the total adds up the rounded figures
    formulas = openpyxl.load_workbook(path)["Jobcard"]
    assert formulas.cell(head + 3, 11).value.startswith("=IF(") and formulas.cell(head + 11, 11).value == f"=SUM(K{head + 1}:K{head + 10})"


def test_stored_results_match_what_the_editor_would_total(repo, tmp_path):
    jc, wp = make(repo, [mine("A", qty=4, norm_hours=1.25), mine("B", qty=0.5, norm_hours=8, factor=2, factor_reason="x")])
    path = tmp_path / "out.xlsx"
    export.write_jobcard(path, jc, wp)
    ws, head = cached(path)
    assert ws.cell(head + 11, 11).value == pytest.approx(jc["totals"]["total"])
    assert ws.cell(head + 3, 11).value == pytest.approx(0.5 * 8 * 2)


def test_a_long_list_stores_every_result_and_the_total(repo, tmp_path):
    jc, wp = make(repo, [mine(f"task {i}", qty=1, norm_hours=i) for i in range(1, 14)])
    path = tmp_path / "out.xlsx"
    export.write_jobcard(path, jc, wp)
    ws, head = cached(path)
    assert [ws.cell(head + i, 11).value for i in range(2, 15)] == list(range(1, 14))
    assert ws.cell(head + 16, 11).value == jc["totals"]["total"] == 1 + sum(range(1, 14)) + 1


def test_storing_results_leaves_a_workbook_that_still_opens_normally(repo, tmp_path):
    jc, wp = make(repo, [mine("A", note="x"), mine("", kind="note", note="<b>b</b>")])
    path = tmp_path / "out.xlsx"
    export.write_jobcard(path, jc, wp)
    assert not list(tmp_path.glob("*.tmp"))                                           # no temporary file left behind
    wb = openpyxl.load_workbook(path, rich_text=True)
    assert wb.sheetnames == ["Jobcard"] and wb["Jobcard"]["C2"].value == "NMC-BB18-001"
    assert wb["Jobcard"].page_setup.fitToWidth == 1 and wb["Jobcard"].row_breaks.brk
