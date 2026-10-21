"""The CTR pages and endpoints, through the real app."""
import io
import sqlite3
import zipfile

import openpyxl
import pytest

from hosts.standalone import create_app

PAYLOAD = {
    "title": "Construction - Piping", "ctr_no": "CTR-014", "client": "Apache",
    "lines": [{"description": "P/Fitter", "grp": "P/Fitter", "eng_override": "137"},
              {"description": "Isometrics", "flag": "D", "qty": "17", "eng_norm": "3", "des_norm": "15"}],
    "resources": [{"name": "Lead", "source": "group:P/Fitter", "split_pct": "60", "rate": "55.5"},
                  {"name": "Senior", "source": "group:P/Fitter", "split_pct": "40", "rate": "48"}],
    "additions": [{"name": "Offshore", "unit_cost": "100", "units": "3"}]}


@pytest.fixture
def app(tmp_path):
    return create_app(tmp_path, "tester")


def new_ctr(client, title="Construction - Piping", ctr_no="CTR-014", workbook="Beryl Bravo diesel mods", **details):
    """A new workbook with one CTR in it; returns the CTR's id. (The first workbook is 1 and its first CTR is 1.)"""
    r = client.post("/ctrs/", data={"title": workbook, **details})
    assert r.status_code == 302 and "/ctrs/estimates/" in r.headers["Location"]
    r = client.post(r.headers["Location"].rsplit("?", 1)[0] + "/ctrs", data={"title": title, "ctr_no": ctr_no})
    assert r.status_code == 302
    return int(r.headers["Location"].rsplit("/", 1)[1])


@pytest.fixture
def c(app):
    client = app.test_client()
    assert new_ctr(client) == 1
    return client


def save(c, body=None, ctr_id=1):
    return c.post(f"/ctrs/{ctr_id}/save", json=PAYLOAD if body is None else body)


def jobcards_snapshot(app):
    """A jobcards workpack with a welder card and a rigger card, and a snapshot of it. Returns the test client."""
    j = app.test_client()
    j.post("/jobcards/workpacks", data={"number": "WP1", "title": "Wellbay"})
    for trade in ("Welder", "Rigger"):
        j.post("/jobcards/workpacks/1/jobcards", data={"trade": trade})
    for jc_id, trade, qty in ((1, "Welder", 10), (2, "Rigger", 5)):
        j.post(f"/jobcards/cards/{jc_id}/save", json={"title": "J", "trade": trade, "tasks": [
            {"description": "work", "qty": qty, "norm_hours": 1, "basis": "q"}]})
    j.post("/jobcards/workpacks/1/snapshots", data={"label": "Issued for CTR rev A"})
    return j


# ---- the list ---------------------------------------------------------------------------------------------------
def test_the_menu_has_ctrs_after_the_other_modules(app):
    assert [n["label"] for n in app.config["NAV"]][-1] == "CTRs"


def test_the_list_starts_empty_and_creating_a_workbook_opens_its_page(app):
    c = app.test_client()
    page = c.get("/ctrs/")
    assert page.status_code == 200 and b"No workbooks yet" in page.data
    r = c.post("/ctrs/", data={"title": "Beryl Bravo", "client": "Apache", "project_name": "BB"})
    assert r.status_code == 302 and r.headers["Location"].endswith("/ctrs/estimates/1")
    listed = c.get("/ctrs/").data.decode()
    assert "Beryl Bravo" in listed and "Apache" in listed and "No workbooks yet" not in listed
    assert b"No CTR sheets yet" in c.get("/ctrs/estimates/1").data


def test_ticking_disciplines_starts_the_workbook_with_a_sheet_for_each(app):
    c = app.test_client()
    r = c.post("/ctrs/", data={"title": "BB", "client": "Apache", "disciplines": ["Piping", "Mechanical", "Process"]})
    page = c.get(r.headers["Location"]).data.decode()
    assert page.index("Process") < page.index("Piping") < page.index("Mechanical")        # the usual order, not the order ticked
    assert page.count('/ctrs/1"') == 1 and page.count('/ctrs/2"') == 1 and page.count('/ctrs/3"') == 1
    assert "Apache" in c.get("/ctrs/1").data.decode()                                     # each sheet has the project details


def test_an_unknown_discipline_is_refused_and_nothing_is_made(app):
    c = app.test_client()
    r = c.post("/ctrs/", data={"title": "BB", "disciplines": ["Piping", "Basket weaving"]})
    assert "error=" in r.headers["Location"] and b"Unknown discipline" in c.get(r.headers["Location"]).data
    assert b"No workbooks yet" in c.get("/ctrs/").data


def test_a_workbook_without_a_name_is_refused_with_a_message(app):
    c = app.test_client()
    r = c.post("/ctrs/", data={"title": "  "})
    assert r.status_code == 302 and "error=" in r.headers["Location"]
    assert b"Estimate name is required" in c.get(r.headers["Location"]).data
    assert b"No workbooks yet" in c.get("/ctrs/").data


def test_a_ctr_without_a_title_is_refused_with_a_message(c):
    r = c.post("/ctrs/estimates/1/ctrs", data={"title": " "})
    assert r.status_code == 302 and "error=" in r.headers["Location"]
    assert b"CTR title is required" in c.get(r.headers["Location"]).data


def test_the_lists_show_hours_and_cost_for_each_ctr_and_the_whole_workbook(c):
    save(c)
    cost = f"£{82 * 55.5 + 55 * 48 + 300:,.2f}"
    for url in ("/ctrs/", "/ctrs/estimates/1"):
        page = c.get(url).data.decode()
        assert "137" in page and cost in page, url


def test_a_workbook_page_has_its_sheets_in_order_and_the_totals(c):
    save(c)
    new_ctr_in_same_workbook = c.post("/ctrs/estimates/1/ctrs", data={"title": "Mechanical"})
    assert new_ctr_in_same_workbook.headers["Location"].endswith("/ctrs/2")
    page = c.get("/ctrs/estimates/1").data.decode()
    assert page.index("Construction - Piping") < page.index("Mechanical")
    assert 'href="/ctrs/estimates/1/export.xlsx"' in page and "Download workbook" in page
    assert c.get("/ctrs/estimates/99").status_code == 404 and c.post("/ctrs/estimates/99/ctrs", data={"title": "x"}).status_code == 404


def test_the_project_details_are_kept_once_and_copied_to_every_sheet(c):
    c.post("/ctrs/estimates/1/ctrs", data={"title": "Mechanical"})
    r = c.post("/ctrs/estimates/1", data={"title": "BB diesel", "client": "Apache", "project_name": "Beryl Bravo",
                                          "project_title": "Diesel mods", "framework_ref": "FA-9"})
    assert r.headers["Location"].endswith("/ctrs/estimates/1?saved=1")
    for pk in (1, 2):
        page = c.get(f"/ctrs/{pk}").data.decode()
        assert 'value="Apache"' in page and 'value="Beryl Bravo"' in page and 'value="FA-9"' in page
        assert 'data-field="client"' not in page and 'data-field="project_name"' not in page      # changed on the workbook, not here
    assert "BB diesel" in c.get("/ctrs/1").data.decode()                                         # the breadcrumb
    assert b"Estimate name is required" in c.get(c.post("/ctrs/estimates/1", data={"title": " "}).headers["Location"]).data


def test_deleting_a_workbook_deletes_its_sheets(c):
    c.post("/ctrs/estimates/1/ctrs", data={"title": "Mechanical"})
    r = c.post("/ctrs/estimates/1/delete")
    assert r.status_code == 302 and c.get("/ctrs/1").status_code == 404 and c.get("/ctrs/2").status_code == 404
    assert c.get("/ctrs/estimates/1").status_code == 404 and b"No workbooks yet" in c.get("/ctrs/").data


# ---- the editor -----------------------------------------------------------------------------------------------
def test_the_editor_page_carries_the_ctr_and_its_endpoints(c):
    save(c)
    page = c.get("/ctrs/1").data.decode()
    for needle in ("Construction - Piping", 'data-save-url="/ctrs/1/save"', 'data-preview-url="/ctrs/1/preview"',
                   'id="ctr-data"', "ctrs.js", "ctrs.css", "Download Excel", "Add allowance", "Resources", "Additions"):
        assert needle in page, needle
    assert c.get("/ctrs/99").status_code == 404 and c.get("/ctrs/static/ctrs.js").status_code == 200
    assert c.get("/ctrs/static/ctrs.css").status_code == 200


def test_the_state_given_to_the_script_holds_the_rows_and_the_worked_out_figures(c):
    import json, re
    save(c)
    page = c.get("/ctrs/1").data.decode()
    state = json.loads(re.search(r'<script type="application/json" id="ctr-data">(.*?)</script>', page, re.S).group(1))
    assert [l["description"] for l in state["lines"]] == ["P/Fitter", "Isometrics"]
    assert [r["hours"] for r in state["computed"]["resources"]] == [82, 55] and state["computed"]["totals"]["hours"] == 137
    assert state["flags"] == ["", "A", "D"]


def test_whatever_is_typed_cannot_break_out_of_the_page(c):
    evil = "</script><script>alert(1)</script><img src=x onerror=alert(2)>"
    r = save(c, {"title": evil, "objectives": evil, "lines": [{"description": evil, "grp": evil}],
                 "resources": [{"name": evil, "source": "pool1"}], "additions": [{"name": evil}]})
    assert r.status_code == 200
    page = c.get("/ctrs/1").data.decode()
    assert "<script>alert" not in page and "<img src=x" not in page and "</script><script>" not in page
    assert c.get("/ctrs/").data.decode().count("<script>alert") == 0 and "<img src=x" not in c.get("/ctrs/").data.decode()


def test_a_viewer_sees_the_page_without_the_controls_to_change_it(app, c):
    app.config["CTX"]._permissions = {"ctrs.view"}
    page = c.get("/ctrs/1").data.decode()
    assert 'data-can-edit="0"' in page and 'id="save"' not in page and 'id="add-line"' not in page
    assert "Delete this CTR" not in page and "Download Excel" in page and "readonly" in page


# ---- saving and previewing ------------------------------------------------------------------------------------
def test_save_returns_what_was_kept_and_the_worked_out_figures(c):
    r = save(c)
    body = r.get_json()
    assert r.status_code == 200 and body["ok"] and body["updated_by"] == "tester"
    assert [l["description"] for l in body["lines"]] == ["P/Fitter", "Isometrics"]
    assert [x["hours"] for x in body["computed"]["resources"]] == [82, 55]
    assert body["computed"]["totals"]["cost"] == round(82 * 55.5 + 55 * 48 + 300, 2)
    assert c.get("/ctrs/1").status_code == 200 and b"Isometrics" in c.get("/ctrs/1").data


def test_a_save_with_a_mistake_says_where_and_changes_nothing(c):
    save(c)
    bad = dict(PAYLOAD, lines=[{"description": "ok"}, {"qty": "lots"}])
    r = save(c, bad)
    assert r.status_code == 400 and r.get_json() == {"error": "Scope line 2: Quantity must be a number"}
    after = save(c).get_json()                                  # still what it was
    assert [l["description"] for l in after["lines"]] == ["P/Fitter", "Isometrics"]


@pytest.mark.parametrize("body", [None, "text", [1, 2], 5])
def test_a_save_that_is_not_a_json_object_is_refused_and_wipes_nothing(c, body):
    save(c)
    r = c.post("/ctrs/1/save", json=body) if body is not None else c.post("/ctrs/1/save", data="not json", content_type="text/plain")
    assert r.status_code == 400 and r.get_json()["error"] == "Send the CTR as JSON"
    assert len(save(c).get_json()["lines"]) == 2
    p = c.post("/ctrs/1/preview", data="x", content_type="text/plain")
    assert p.status_code == 400


def test_a_save_to_a_missing_ctr_is_a_404(c):
    assert save(c, ctr_id=99).status_code == 404
    assert c.post("/ctrs/99/preview", json={}).status_code == 404


def test_a_preview_returns_figures_for_what_is_on_the_page_and_saves_nothing(c):
    save(c)
    r = c.post("/ctrs/1/preview", json=dict(PAYLOAD, resources=[
        {"name": "Solo", "source": "group:P/Fitter", "split_pct": "100", "rate": "10"}, {}]))
    body = r.get_json()
    assert r.status_code == 200 and body["ok"]
    assert [x["hours"] for x in body["computed"]["resources"]] == [137, 0]          # one answer per row sent, blank included
    saved = save(c).get_json()
    assert [x["name"] for x in saved["resources"]] == ["Lead", "Senior"]            # the preview changed nothing


def test_a_preview_names_the_mistake_it_finds(c):
    r = c.post("/ctrs/1/preview", json={"resources": [{"name": "A", "split_pct": "150"}]})
    assert r.status_code == 400 and r.get_json() == {"error": "Resource 1: Split cannot be above 100"}


def test_a_preview_reports_the_checks(c):
    r = c.post("/ctrs/1/preview", json={"lines": [{"description": "Weld", "eng_override": "10"}], "resources": []})
    assert r.get_json()["computed"]["warnings"] == ["No resource takes the Eng hours of: Weld."]


# ---- importing hours from the jobcards ---------------------------------------------------------------------
def test_the_snapshots_to_choose_from_are_listed_with_their_workpack_and_hours(app, c):
    page = c.get("/ctrs/1").data.decode()
    assert "Hours from the jobcards" in page and "No snapshots yet" in page and "Import hours" not in page
    jobcards_snapshot(app)
    page = c.get("/ctrs/1").data.decode()
    assert "WP1 — Issued for CTR rev A" in page and "Import hours" in page and "No snapshots yet" not in page
    assert "19 h" in page                                       # 10 + 5 + two standard tasks on each of two cards


def test_importing_brings_in_a_line_per_trade_and_says_so(app, c):
    jobcards_snapshot(app)
    r = c.post("/ctrs/1/import-hours", data={"snapshot_id": "1"})
    assert r.status_code == 302 and "imported=2" in r.headers["Location"] and r.headers["Location"].endswith("#scope")
    page = c.get(r.headers["Location"].split("#")[0]).data.decode()
    assert "2 trade lines brought in from the jobcards" in page and "Built from: <b>WP1 - Issued for CTR rev A" in page
    lines = save(c, {"title": "x"}).get_json()["lines"]
    assert [(l["description"], l["grp"], l["eng_override"], l["from_snapshot"]) for l in lines] == [
        ("Rigger", "Rigger", 7, 1), ("Welder", "Welder", 12, 1)]                # each card also has the two standard tasks


def test_an_imported_ctr_can_be_set_up_end_to_end_with_splits(app, c):
    jobcards_snapshot(app)
    c.post("/ctrs/1/import-hours", data={"snapshot_id": "1"})
    body = {"title": "Construction - Piping", "resources": [
        {"name": "Lead Welder", "source": "group:Welder", "split_pct": "60", "rate": "50"},
        {"name": "Senior Welder", "source": "group:Welder", "split_pct": "40", "rate": "45"},
        {"name": "Rigger", "source": "group:Rigger", "split_pct": "100", "rate": "40"}]}
    out = save(c, body).get_json()["computed"]
    assert [r["hours"] for r in out["resources"]] == [7, 5, 7]                   # 12 hours at 60/40 -> 7.2 and 4.8 -> 7 + 5
    assert out["totals"]["hours"] == 19 and out["warnings"] == []               # all of the jobcards' 19 hours are taken once


def test_importing_again_after_the_jobcards_change_replaces_the_old_import(app, c):
    j = jobcards_snapshot(app)
    c.post("/ctrs/1/import-hours", data={"snapshot_id": "1"})
    j.post("/jobcards/cards/2/save", json={"title": "J", "trade": "Rigger", "tasks": [{"description": "w", "qty": 50, "norm_hours": 1, "basis": "q"}]})
    j.post("/jobcards/workpacks/1/snapshots", data={"label": "rev B"})
    c.post("/ctrs/1/import-hours", data={"snapshot_id": "2"})
    lines = save(c, {"title": "x"}).get_json()["lines"]
    assert [(l["description"], l["eng_override"]) for l in lines] == [("Rigger", 52), ("Welder", 12)]
    assert "rev B" in c.get("/ctrs/1").data.decode().split("Built from:")[1][:80]


@pytest.mark.parametrize("form", [{}, {"snapshot_id": ""}, {"snapshot_id": "abc"}, {"snapshot_id": "99"}, {"snapshot_id": "1.5"}])
def test_importing_nothing_or_an_unknown_snapshot_is_refused_with_a_message(app, c, form):
    jobcards_snapshot(app)
    r = c.post("/ctrs/1/import-hours", data=form)
    assert r.status_code == 302 and "error=" in r.headers["Location"]
    page = c.get(r.headers["Location"]).data.decode()
    assert ("Choose a snapshot to import" in page or "No such snapshot" in page)
    assert save(c, {"title": "x"}).get_json()["lines"] == []


def test_a_snapshot_with_no_hours_is_refused(app, c):
    j = app.test_client()
    j.post("/jobcards/workpacks", data={"number": "EMPTY"})
    j.post("/jobcards/workpacks/1/snapshots", data={"label": "nothing"})
    r = c.post("/ctrs/1/import-hours", data={"snapshot_id": "1"})
    assert "error=" in r.headers["Location"] and b"no hours" in c.get(r.headers["Location"]).data


def test_without_the_jobcards_module_the_ctr_module_still_works_and_offers_no_import(tmp_path):
    app = create_app(tmp_path, "tester", modules=["ctrs"])
    c = app.test_client()
    assert new_ctr(c, "Alone") == 1
    page = c.get("/ctrs/1").data.decode()
    assert page.count("Hours from the jobcards") == 0 and "import-form" not in page and "snapshot_id" not in page
    assert save(c, {"title": "Alone", "lines": [{"description": "x", "eng_override": "5"}]}).status_code == 200
    r = c.post("/ctrs/1/import-hours", data={"snapshot_id": "1"})
    assert b"jobcards module is not installed" in c.get(r.headers["Location"]).data
    assert c.get("/ctrs/1/export.xlsx").status_code == 200


# ---- duplicate and delete ------------------------------------------------------------------------------------
def test_duplicate_makes_a_copy_and_opens_it(c):
    save(c)
    r = c.post("/ctrs/1/duplicate")
    assert r.status_code == 302 and r.headers["Location"].endswith("/ctrs/2")
    assert b"Construction - Piping (copy)" in c.get("/ctrs/2").data and b"Isometrics" in c.get("/ctrs/2").data
    assert c.post("/ctrs/99/duplicate").status_code == 404


def test_delete_removes_the_ctr_and_goes_back_to_its_workbook(c):
    r = c.post("/ctrs/1/delete")
    assert r.status_code == 302 and r.headers["Location"].endswith("/ctrs/estimates/1") and c.get("/ctrs/1").status_code == 404
    assert c.post("/ctrs/1/delete").status_code == 404


# ---- the Excel file ------------------------------------------------------------------------------------------
def test_the_download_is_an_xlsx_named_after_the_ctr_that_opens(c):
    save(c)
    r = c.get("/ctrs/1/export.xlsx")
    assert r.status_code == 200 and r.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "attachment" in r.headers["Content-Disposition"] and "CTR-014 Construction - Piping.xlsx" in r.headers["Content-Disposition"]
    ws = openpyxl.load_workbook(io.BytesIO(r.data)).active
    assert ws["B11"].value == "Construction - Piping" and ws["I11"].value == "CTR-014"
    assert zipfile.ZipFile(io.BytesIO(r.data)).testzip() is None


def download_name(c):
    import re
    return re.search(r'filename="?([^";]+)', c.get("/ctrs/1/export.xlsx").headers["Content-Disposition"]).group(1)


def test_the_file_name_cannot_carry_path_characters(c):
    save(c, {"title": r'A/B\C:D*E?"F<G>H|I', "ctr_no": ""})
    assert download_name(c) == "ABCDEFGHI.xlsx"
    save(c, {"title": "x", "ctr_no": "///"})
    assert download_name(c) == "x.xlsx"
    save(c, {"title": "///", "ctr_no": ""})
    assert download_name(c) == "CTR.xlsx"
    save(c, {"title": "Piping", "ctr_no": "CTR 7"})
    assert download_name(c) == "CTR 7 Piping.xlsx"


def test_the_download_is_made_from_what_is_saved(c):
    save(c)
    save(c, dict(PAYLOAD, objectives="Saved objectives"))
    ws = openpyxl.load_workbook(io.BytesIO(c.get("/ctrs/1/export.xlsx").data)).active
    assert ws["A16"].value == "Saved objectives"


def test_a_missing_ctr_has_no_download(c):
    assert c.get("/ctrs/99/export.xlsx").status_code == 404


def test_the_download_leaves_no_temporary_files_behind(c, tmp_path):
    save(c)
    before = {p.name for p in tmp_path.rglob("*")}
    c.get("/ctrs/1/export.xlsx")
    assert {p.name for p in tmp_path.rglob("*")} == before


# ---- permissions, safety and the audit log ---------------------------------------------------------------------
@pytest.mark.parametrize("method, url, kwargs", [
    ("post", "/ctrs/", {"data": {"title": "x"}}), ("post", "/ctrs/1/save", {"json": PAYLOAD}),
    ("post", "/ctrs/1/import-hours", {"data": {"snapshot_id": "1"}}), ("post", "/ctrs/1/duplicate", {}), ("post", "/ctrs/1/delete", {})])
def test_a_viewer_cannot_change_anything(app, c, method, url, kwargs):
    app.config["CTX"]._permissions = {"ctrs.view"}
    assert getattr(c, method)(url, **kwargs).status_code == 403
    assert c.get("/ctrs/1").status_code == 200 and c.get("/ctrs/1/export.xlsx").status_code == 200
    assert c.post("/ctrs/1/preview", json={}).status_code == 200


def test_someone_with_no_permissions_sees_nothing(app, c):
    app.config["CTX"]._permissions = set()
    for url in ("/ctrs/", "/ctrs/1", "/ctrs/1/export.xlsx"):
        assert c.get(url).status_code == 403, url
    assert c.post("/ctrs/1/preview", json={}).status_code == 403


@pytest.mark.parametrize("headers", [{"Origin": "http://evil.example"}, {"Sec-Fetch-Site": "cross-site"}, {"Referer": "http://evil.example/x"}])
def test_requests_from_other_sites_are_refused(c, headers):
    save(c)
    for url, kwargs in (("/ctrs/", {"data": {"title": "x"}}), ("/ctrs/1/save", {"json": PAYLOAD}), ("/ctrs/1/delete", {}),
                        ("/ctrs/1/import-hours", {"data": {"snapshot_id": "1"}}), ("/ctrs/1/preview", {"json": {}})):
        assert c.post(url, headers=headers, **kwargs).status_code == 403, url
    assert c.get("/ctrs/1").status_code == 200


def test_changes_are_written_to_the_audit_log(app, c, tmp_path):
    save(c)
    c.get("/ctrs/1/export.xlsx")
    c.post("/ctrs/1/duplicate")
    c.post("/ctrs/2/delete")
    c.get("/ctrs/estimates/1/export.xlsx")
    c.post("/ctrs/estimates/1", data={"title": "Renamed", "client": "Apache"})
    c.post("/ctrs/estimates/1/delete")
    with sqlite3.connect(tmp_path / "audit.sqlite") as db:
        rows = db.execute("SELECT user, action, detail FROM audit WHERE module='ctrs' ORDER BY id").fetchall()
    assert [r[1] for r in rows] == ["estimate_created", "ctr_created", "ctr_saved", "ctr_downloaded", "ctr_created", "ctr_saved",
                                    "ctr_duplicated", "ctr_deleted", "estimate_downloaded", "estimate_saved", "estimate_deleted"]
    assert {r[0] for r in rows} == {"tester"} and ("tester", "ctr_downloaded", "Construction - Piping") in rows
    assert ("tester", "estimate_downloaded", "Beryl Bravo diesel mods") in rows and ("tester", "estimate_deleted", "Renamed") in rows


def test_a_preview_is_not_audited(c, tmp_path):
    n = lambda: sqlite3.connect(tmp_path / "audit.sqlite").execute("SELECT COUNT(*) FROM audit").fetchone()[0]
    before = n()
    c.post("/ctrs/1/preview", json=PAYLOAD)
    assert n() == before


# ---- the look and the script ----------------------------------------------------------------------------------
def test_the_stylesheet_has_sharp_edges_and_the_house_colours(c):
    css = c.get("/ctrs/static/ctrs.css").data.decode()
    assert "border-radius" not in css and "var(--accent)" in css and "var(--ink)" in css


# ---- the friendlier editor ---------------------------------------------------------------------------------------
def test_the_editor_leads_with_the_steps_and_the_checks_and_has_a_jump_menu(c):
    page = c.get("/ctrs/1").data.decode()
    for step in ("hours", "resources", "rates", "download"):
        assert f'data-step="{step}"' in page
    assert page.index('id="steps"') < page.index('id="checks"') < page.index('class="jump"') < page.index('id="scope"')
    assert page.index('id="warnings"') < page.index('id="lines"')                     # the checks are above the tables
    for target in ("sec-header", "hours-source", "scope", "sec-notes", "resources-h", "additions-h", "sec-terms"):
        assert f'href="#{target}"' in page and f'id="{target}"' in page, target
    assert 'id="checks-badge"' in page


def test_the_long_text_sections_are_folded_and_the_header_opens_while_it_is_unfinished(c):
    import re
    page = c.get("/ctrs/1").data.decode()
    assert re.search(r'<details class="sec" id="sec-header" open>', page)              # no objectives yet: left open
    assert re.search(r'<details class="sec" id="sec-notes">', page) and re.search(r'<details class="sec" id="sec-terms">', page)
    save(c, {"title": "x", "objectives": "Done"})
    assert re.search(r'<details class="sec" id="sec-header" >', c.get("/ctrs/1").data.decode())   # filled in: folded away


def test_the_wording_is_plain_and_there_is_a_button_for_a_resource_per_trade(c):
    page = c.get("/ctrs/1").data.decode()
    assert "Hours come from" in page and "Trade / group" in page and "Add a resource for each trade" in page
    assert "Takes hours from" not in page and ">Group<" not in page


def test_a_viewer_does_not_get_the_buttons_that_add_rows(app, c):
    app.config["CTX"]._permissions = {"ctrs.view"}
    page = c.get("/ctrs/1").data.decode()
    assert "add-trades" not in page and "add-resource" not in page


def test_the_preview_says_how_far_along_the_ctr_is_and_the_split_totals(c):
    body = c.post("/ctrs/1/preview", json={"lines": [{"description": "x", "grp": "G", "eng_override": "10"}],
                                          "resources": [{"name": "A", "source": "group:G", "split_pct": "60", "rate": "1"},
                                                        {"name": "B", "source": "group:G", "split_pct": "40"}]}).get_json()["computed"]
    assert body["steps"] == {"hours": True, "resources": True, "rates": False} and body["split_totals"] == {"group:G": 100}


def test_a_new_ctr_comes_with_who_made_it_and_the_date(app):
    from datetime import date
    c = app.test_client()
    new_ctr(c, "Fresh")
    page = c.get("/ctrs/1").data.decode()
    assert 'data-field="prepared_by" ' in page and 'value="tester"' in page and f'value="{date.today().isoformat()}"' in page


def test_the_script_asks_before_downloading_with_open_checks_and_helps_with_splits(c):
    js = c.get("/ctrs/static/ctrs.js").data.decode()
    for hook in ("Download anyway?", "completeSplit", "addTradeResources", "split_totals", "c.steps", "checks-badge", ".jump a"):
        assert hook in js, hook


def test_the_script_does_no_sums_of_its_own_and_uses_no_unsafe_html(c):
    js = c.get("/ctrs/static/ctrs.js").data.decode()
    assert "innerHTML" not in js and "eval(" not in js and "document.write" not in js
    assert "Math.ceil" not in js and "Math.round" not in js and "Math.floor" not in js       # the server does all the sums
    for hook in ("previewUrl", "saveUrl", "beforeunload", "setTimeout(preview"):
        assert hook in js, hook
